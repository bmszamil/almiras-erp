"""The fillable migration workbook (migration.md §11): generator (W1), reader (W3) and adapters to the loaders.

``build_workbook`` writes the template from ``workbook_spec``; ``read_workbook`` parses a filled copy into typed
records per sheet without touching the database; ``trial_balance_for`` and ``account_map_rows_for`` turn one
entity's sheets into the inputs of the existing services (``validate_trial_balance``, ``load_opening_balances``,
``create_account_map``). Checks across sheets live in ``workbook_checks``.

A filled workbook holds legacy and personal data: read it from ``exports/`` (ignored by git), never commit it (W2).
Problems never repeat the value of a personal-data column.
"""
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from banking import iban as iban_rules

from . import workbook_spec as spec
from .readers import (CENT, ZERO, CellError, ParsedFile, Problem, ReaderError, TBRow, bi, normalise_header, to_amount,
                      to_code, to_date, to_int, to_text)

DEFAULT_PATH = Path(__file__).resolve().parent.parent / "docs" / "migration" / "almiras-migration-workbook.xlsx"
GENERATED_AT = datetime(2026, 10, 18)  # fixed document properties: regenerating gives the same content
HEADER_SCAN_ROWS = 20
_ENTITY_CODE = re.compile(r"^[A-Z0-9]{2,8}$")
_ERP_CODE = re.compile(r"^\d{3,8}$")
_DIGITS = re.compile(r"^\d+$")

# colours
FILL_TITLE, FILL_HEADER, FILL_REQUIRED, FILL_PERSONAL = "1F4E78", "DDEBF7", "F8CBAD", "D9C3E9"
FILL_KEY, FILL_EXAMPLE = "F2F2F2", "EDEDED"


class WorkbookProblem(Problem):
    """A finding with the sheet and entity it belongs to."""

    def __init__(self, code, severity, message, row=None, entry_no="", account_code="", sheet="", entity=""):
        super().__init__(code, severity, message, row, entry_no, account_code)
        self.sheet, self.entity = sheet, entity

    def as_dict(self):
        return {"sheet": self.sheet, "entity": self.entity, **super().as_dict()}


def problem(severity, code, en, ar, *, sheet="", row=None, entity="", account_code=""):
    return WorkbookProblem(code, severity, bi(en, ar), row=row, account_code=account_code, sheet=sheet, entity=entity)


# --- generator (W1) ------------------------------------------------------------------------------------------------

def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFBFBF")
    return {
        "title": Font(bold=True, size=14, color="FFFFFF"), "title_fill": PatternFill("solid", fgColor=FILL_TITLE),
        "bold": Font(bold=True), "key": Font(size=8, color="7F7F7F"), "example": Font(italic=True, color="7F7F7F"),
        "fill": {k: PatternFill("solid", fgColor=c) for k, c in (("header", FILL_HEADER), ("required", FILL_REQUIRED),
                                                                 ("personal", FILL_PERSONAL), ("key", FILL_KEY),
                                                                 ("example", FILL_EXAMPLE))},
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "rtl": Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
    }


NUMBER_FORMATS = {"sar": "#,##0.00", "amount": "#,##0.000", "rate": "0.##", "date": "yyyy-mm-dd", "int": "0"}
TEXT_KINDS = {"entity", "text", "legacy_code", "erp_code", "vat", "cr", "iban", "currency", "country", "choice", "yesno"}
LIST_SHEET = "Lists"


def list_name(key):
    return f"L_{key}"


def _list_columns():
    """(key, header, values) of the Lists sheet: dropdown labels, currencies, countries, ERP codes."""
    from ledger.coa_template import TEMPLATE

    out = [(k, k, [spec.label(c) for c in choices]) for k, choices in spec.CHOICES.items()]
    out.append(("currency", "currency", sorted(spec.CURRENCIES)))
    out.append(("country", "country", spec.COUNTRIES))
    out.append(("erp_code", "erp_code", [code for code, _en, _ar, _t, flags in TEMPLATE
                                         if not flags.get("header") and not flags.get("report_only")]))
    return out


def _validation(column, sheet_name):
    """The DataValidation for a column (None when the column has none)."""
    from openpyxl.worksheet.datavalidation import DataValidation

    kw = dict(allow_blank=True, showErrorMessage=True, showInputMessage=True)
    if column.kind in ("choice", "yesno", "currency", "country", "entity", "erp_code"):
        source = {"choice": column.choices, "yesno": "yes_no"}.get(column.kind, column.kind)
        dv = DataValidation(type="list", formula1=list_name(source), **kw)
        dv.errorTitle = "Choose from the list"
        dv.error = "Choose a value from the dropdown list. / اختر قيمة من القائمة المنسدلة."
        if column.kind in ("entity", "erp_code"):  # an entity's own ERP chart may have more codes than the template
            dv.errorStyle = "warning"
            dv.error = ("Not in the list: check the code. / القيمة ليست في القائمة: تحقق من الرمز." if column.kind ==
                        "erp_code" else "Add the entity on the Entities sheet first. / أضف الكيان في ورقة الكيانات أولًا.")
        return dv
    if column.kind in ("sar", "amount", "rate"):
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", **kw)
        dv.errorTitle = "Amount"
        dv.error = ("A number of 0 or more, no text. / رقم يساوي صفرًا أو أكثر، دون نص." if column.kind != "sar" else
                    "SAR, 0 or more, two decimals at most. / بالريال، صفر أو أكثر، بخانتين عشريتين كحد أقصى.")
        return dv
    if column.kind == "date":
        dv = DataValidation(type="date", operator="between", formula1="DATE(1950,1,1)", formula2="DATE(2100,12,31)", **kw)
        dv.errorTitle = "Date"
        dv.error = "A date, e.g. 2026-12-31. / تاريخ، مثل 2026-12-31."
        return dv
    if column.kind == "int":
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", **kw)
        dv.errorTitle = "Whole number"
        dv.error = "A whole number of 0 or more. / عدد صحيح يساوي صفرًا أو أكثر."
        return dv
    if column.kind == "vat":
        dv = DataValidation(type="textLength", operator="equal", formula1="15", **kw)
        dv.errorTitle = "VAT number"
        dv.error = "15 digits, starting and ending with 3. / 15 رقمًا تبدأ وتنتهي بالرقم 3."
        return dv
    if column.kind == "iban":
        dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="34", **kw)
        dv.errorTitle = "IBAN"
        dv.error = "An IBAN has at most 34 characters. / الآيبان 34 خانة كحد أقصى."
        return dv
    return None


def _write_data_sheet(wb, sheet, st):
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(sheet.name)
    last = get_column_letter(len(sheet.columns))
    ws["A1"] = f"{sheet.name} / {sheet.title_ar}"
    ws["A1"].font, ws["A1"].fill = st["title"], st["title_fill"]
    ws.merge_cells(f"A1:{last}1")
    ws["A2"] = (f"{sheet.intro_en}  * = required  [PD] = personal data  Grey rows from row 6 are fake examples: delete "
                f"them. Do not change rows 1-5.\n{sheet.intro_ar}  * = مطلوب  [PD] = بيانات شخصية  الصفوف الرمادية "
                "من الصف 6 أمثلة وهمية: احذفها. لا تغيّر الصفوف 1-5.")
    ws["A2"].alignment = st["wrap"]
    ws.merge_cells(f"A2:{last}2")
    ws.row_dimensions[2].height = 48
    ws.row_dimensions[spec.EN_ROW].height = 60
    ws.row_dimensions[spec.AR_ROW].height = 60
    last_row = spec.FIRST_DATA_ROW + spec.MAX_ROWS - 1
    for i, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(i)
        fill = st["fill"]["personal" if column.personal else "required" if column.required else "header"]
        for row, text, align in ((spec.EN_ROW, column.header_en, st["wrap"]), (spec.AR_ROW, column.header_ar, st["rtl"])):
            cell = ws.cell(row=row, column=i, value=text)
            cell.font, cell.fill, cell.alignment, cell.border = st["bold"], fill, align, st["border"]
        key = ws.cell(row=spec.KEY_ROW, column=i, value=column.key)
        key.font, key.fill = st["key"], st["fill"]["key"]
        dim = ws.column_dimensions[letter]
        dim.width = column.width
        dim.number_format = NUMBER_FORMATS.get(column.kind, "@")
        dv = _validation(column, sheet.name)
        if dv is not None:
            dv.promptTitle = column.en[:32]
            dv.prompt = f"{column.header_en} / {column.header_ar}"[:250]
            dv.add(f"{letter}{spec.FIRST_DATA_ROW}:{letter}{last_row}")
            ws.add_data_validation(dv)
    for n, example in enumerate(sheet.examples):
        row = spec.FIRST_DATA_ROW + n
        for i, column in enumerate(sheet.columns, start=1):
            value = example.get(column.key)
            if column.kind == "choice" and value is not None:
                value = spec.label(next(c for c in spec.CHOICES[column.choices] if c[0] == value))
            elif column.kind == "yesno" and value is not None:
                value = spec.label(next(c for c in spec.CHOICES["yes_no"] if c[0] == value))
            elif isinstance(value, Decimal):
                value = float(value)  # an xlsx number; read back through its shortest repr (R2)
            cell = ws.cell(row=row, column=i, value=value)
            cell.font, cell.fill = st["example"], st["fill"]["example"]
            cell.number_format = NUMBER_FORMATS.get(column.kind, "@")
    ws.freeze_panes = ws.cell(row=spec.FIRST_DATA_ROW, column=2)
    return ws


def _write_instructions(wb, st):
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 90
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["B1"] = "English", "العربية"
    for cell in (ws["A1"], ws["B1"]):
        cell.font, cell.fill = st["title"], st["title_fill"]
    for n, pair in enumerate(spec.INSTRUCTIONS, start=2):
        if pair is None:
            continue
        en, ar = pair
        heading = en.startswith("#")
        a = ws.cell(row=n, column=1, value=en.lstrip("# ") if heading else en)
        b = ws.cell(row=n, column=2, value=ar.lstrip("# ") if heading else ar)
        a.alignment, b.alignment = st["wrap"], st["rtl"]
        if heading:
            a.font = b.font = st["bold"]
    return ws


def _write_checklist(wb, st):
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet("Checklist")
    ws["A1"] = "Checklist and sign-off / قائمة التحقق والاعتماد"
    ws["A1"].font, ws["A1"].fill = st["title"], st["title_fill"]
    ws.merge_cells(f"A1:{get_column_letter(len(spec.CHECKLIST_COLUMNS))}1")
    for i, (_key, en, ar, width) in enumerate(spec.CHECKLIST_COLUMNS, start=1):
        for row, text, align in ((spec.EN_ROW, en, st["wrap"]), (spec.AR_ROW, ar, st["rtl"])):
            cell = ws.cell(row=row, column=i, value=text)
            cell.font, cell.fill, cell.alignment = st["bold"], st["fill"]["header"], align
        ws.column_dimensions[get_column_letter(i)].width = width
    first = spec.FIRST_DATA_ROW
    for n, (en, ar, fills, approves) in enumerate(spec.CHECKLIST_STEPS):
        row = first + n
        for i, value in enumerate((n + 1, en, ar, fills, approves), start=1):
            cell = ws.cell(row=row, column=i, value=value)
            cell.alignment = st["rtl"] if i == 3 else st["wrap"]
        ws.cell(row=row, column=7).number_format = "yyyy-mm-dd"
    status_col = get_column_letter([k for k, *_ in spec.CHECKLIST_COLUMNS].index("status") + 1)
    last = first + len(spec.CHECKLIST_STEPS) - 1
    dv = DataValidation(type="list", formula1=list_name("checklist_status"), allow_blank=True)
    dv.add(f"{status_col}{first}:{status_col}{last}")
    ws.add_data_validation(dv)
    row = last + 3
    ws.cell(row=row, column=2, value="Sign-off of the tie-out report / اعتماد تقرير المطابقة").font = st["bold"]
    for i, (en, ar, _w) in enumerate((("Role", "الدور", 0), ("Name", "الاسم", 0), ("Date", "التاريخ", 0),
                                      ("Signature", "التوقيع", 0)), start=2):
        cell = ws.cell(row=row + 1, column=i, value=f"{en} / {ar}")
        cell.font, cell.fill = st["bold"], st["fill"]["header"]
    for n, (en, ar) in enumerate(spec.SIGN_OFF, start=row + 2):
        ws.cell(row=n, column=2, value=f"{en} / {ar}").alignment = st["wrap"]
    return ws


def _write_erp_chart(wb, st):
    from ledger.coa_template import TEMPLATE

    from .services import CONTROL_CODES

    ws = wb.create_sheet("ERP chart")
    ws["A1"] = "ERP group chart of accounts (reference) / دليل حسابات المجموعة في النظام (للاطلاع)"
    ws["A1"].font, ws["A1"].fill = st["title"], st["title_fill"]
    ws.merge_cells("A1:H1")
    ws["A2"] = ("The group template; an entity's chart in the ERP may have more accounts. Control accounts: open items "
                "must tie to them. / القالب العام للمجموعة؛ قد يحتوي دليل الكيان في النظام على حسابات إضافية. "
                "الحسابات الرقابية: يجب أن تطابقها الأرصدة المفتوحة.")
    ws["A2"].alignment = st["wrap"]
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 32
    heads = (("Code", "الرمز", 8), ("Name (English)", "الاسم (بالإنجليزية)", 40), ("Name (Arabic)", "الاسم (بالعربية)", 40),
             ("Type", "النوع", 12), ("Postable", "قابل للترحيل", 10), ("Intercompany", "شركات المجموعة", 12),
             ("Sharia-prohibited", "محظور شرعًا", 12), ("Control account", "حساب رقابي", 12))
    from openpyxl.utils import get_column_letter
    for i, (en, ar, width) in enumerate(heads, start=1):
        for row, text in ((spec.EN_ROW, en), (spec.AR_ROW, ar)):
            cell = ws.cell(row=row, column=i, value=text)
            cell.font, cell.fill = st["bold"], st["fill"]["header"]
        ws.column_dimensions[get_column_letter(i)].width = width
    yes = {True: "Yes / نعم", False: "No / لا"}
    for n, (code, en, ar, kind, flags) in enumerate(TEMPLATE, start=spec.FIRST_DATA_ROW):
        values = (code, en, ar, kind, yes[not flags.get("header") and not flags.get("report_only")],
                  yes[bool(flags.get("intercompany"))], yes[bool(flags.get("prohibited"))], yes[code in CONTROL_CODES])
        for i, value in enumerate(values, start=1):
            cell = ws.cell(row=n, column=i, value=value)
            if i == 3:
                cell.alignment = st["rtl"]
            if i == 1:
                cell.number_format = "@"
    ws.freeze_panes = ws.cell(row=spec.FIRST_DATA_ROW, column=2)
    ws.protection.sheet = True  # reference only (no password: protects against slips, not people)
    return ws


def _write_lists(wb, st):
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName

    ws = wb.create_sheet(LIST_SHEET)
    for i, (key, header, values) in enumerate(_list_columns(), start=1):
        letter = get_column_letter(i)
        ws.cell(row=1, column=i, value=header).font = st["bold"]
        for n, value in enumerate(values, start=2):
            ws.cell(row=n, column=i, value=value).number_format = "@"
        ws.column_dimensions[letter].width = 34 if key in spec.CHOICES else 10
        wb.defined_names[list_name(key)] = DefinedName(
            list_name(key), attr_text=f"'{LIST_SHEET}'!${letter}$2:${letter}${len(values) + 1}")
    entities = spec.SHEET_BY_KEY["entities"]
    first, last = spec.FIRST_DATA_ROW, spec.FIRST_DATA_ROW + spec.MAX_ROWS - 1
    wb.defined_names[list_name("entity")] = DefinedName(
        list_name("entity"), attr_text=f"'{entities.name}'!$A${first}:$A${last}")
    ws.protection.sheet = True
    return ws


def build_workbook(path=None):
    """The template as xlsx bytes; written to ``path`` too when given."""
    from openpyxl import Workbook

    st = _styles()
    wb = Workbook()
    _write_instructions(wb, st)
    for sheet in spec.SHEETS:
        _write_data_sheet(wb, sheet, st)
    _write_checklist(wb, st)
    _write_erp_chart(wb, st)
    _write_lists(wb, st)
    wb.properties.creator = "Almiras ERP (manage.py migration_template)"
    wb.properties.title = "Almiras migration workbook / ملف ترحيل بيانات الميراس"
    wb.properties.created = wb.properties.modified = GENERATED_AT
    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    if path is not None:
        Path(path).write_bytes(data)
    return data


# --- reader (W3) ---------------------------------------------------------------------------------------------------

def _choice_lookup():
    out = {}
    for key, choices in spec.CHOICES.items():
        table = {}
        for choice in choices:
            code, en, ar = choice
            for text in (code, en, ar, spec.label(choice), f"{en}/{ar}"):
                table[normalise_header(text)] = code
        out[key] = table
    out["yes_no"].update({normalise_header(t): c for t, c in (("y", "yes"), ("n", "no"), ("true", "yes"),
                                                              ("false", "no"), ("1", "yes"), ("0", "no"))})
    return out


CHOICE_LOOKUP = _choice_lookup()


def _choice(value, key):
    if isinstance(value, bool) and key == "yes_no":
        return "yes" if value else "no"
    code = CHOICE_LOOKUP[key].get(normalise_header(to_text(value)))
    if code is None:
        raise CellError(f"{value!r} is not one of the list values.", f"{value!r} ليست من قيم القائمة.")
    return code


def convert(column, value):
    """The typed value of a cell (None when blank). Raises CellError."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    kind = column.kind
    if kind in ("sar", "amount", "rate"):
        amount = to_amount(value)
        if amount < 0:
            raise CellError(f"{amount} is negative; amounts are entered as positive numbers.",
                            f"{amount} سالب؛ تُدخل المبالغ أرقامًا موجبة.")
        if kind == "sar" and amount != amount.quantize(CENT):
            raise CellError(f"{amount} has more than 2 decimals (SAR).", f"{amount} بأكثر من خانتين عشريتين (ريال).")
        return amount.quantize(CENT) if kind == "sar" else amount
    if kind == "date":
        return to_date(value)
    if kind == "int":
        return to_int(value, column.en, column.ar)
    if kind == "legacy_code":
        return to_code(value)
    if kind == "choice":
        return _choice(value, column.choices)
    if kind == "yesno":
        return _choice(value, "yes_no") == "yes"
    text = to_text(value)
    if kind == "entity":
        text = text.upper()
        if not _ENTITY_CODE.match(text):
            raise CellError(f"Entity code {text!r} is not 2 to 8 capitals or digits.",
                            f"رمز الكيان {text!r} ليس من 2 إلى 8 أحرف كبيرة أو أرقام.")
    elif kind == "erp_code":
        if not _ERP_CODE.match(text):
            raise CellError(f"ERP account code {text!r} is not 3 to 8 digits.", f"رمز الحساب {text!r} ليس من 3 إلى 8 أرقام.")
    elif kind == "currency":
        text = text.upper()
        if text not in spec.CURRENCIES:
            raise CellError(f"{text!r} is not a currency code the ERP knows.", f"{text!r} ليس رمز عملة معروفًا في النظام.")
    elif kind == "country":
        text = text.upper()
        if text not in spec.COUNTRIES:
            raise CellError(f"{text!r} is not an ISO 3166 country code.", f"{text!r} ليس رمز دولة وفق ISO 3166.")
    elif kind == "vat":
        if not (len(text) == 15 and _DIGITS.match(text) and text[0] == "3" and text[-1] == "3"):
            raise CellError("A VAT number has 15 digits and starts and ends with 3.",
                            "يتكون الرقم الضريبي من 15 رقمًا ويبدأ وينتهي بالرقم 3.")
    elif kind == "cr":
        if not _DIGITS.match(text):
            raise CellError("A commercial registration number has digits only.", "رقم السجل التجاري أرقام فقط.")
    elif kind == "iban":
        text = iban_rules.normalise(text)
    if len(text) > column.max_len:
        raise CellError(f"Longer than {column.max_len} characters.", f"أطول من {column.max_len} حرفًا.")
    return text


@dataclass
class Record:
    sheet: str
    row: int
    values: dict
    example: bool = False
    invalid: set = field(default_factory=set)  # keys whose cell could not be read (already reported)

    def __getitem__(self, key):
        return self.values.get(key)

    def get(self, key, default=None):
        value = self.values.get(key)
        return default if value is None else value


@dataclass
class SheetData:
    key: str
    name: str
    present: bool = False
    records: list = field(default_factory=list)  # data rows, examples excluded
    examples: list = field(default_factory=list)
    columns: dict = field(default_factory=dict)  # key -> 0-based column index
    header_row: int | None = None


@dataclass
class WorkbookData:
    name: str
    sha256: str
    sheets: dict
    problems: list = field(default_factory=list)
    other_sheets: list = field(default_factory=list)

    def records(self, key, entity=None):
        rows = self.sheets[key].records
        return [r for r in rows if entity is None or r["entity"] == entity]

    def entity(self, code):
        return next((r for r in self.sheets["entities"].records if r["entity_code"] == code), None)

    def entity_codes(self):
        return [r["entity_code"] for r in self.sheets["entities"].records if r["entity_code"]]


def _labels(column):
    plain = lambda t: normalise_header(t.replace("[PD]", "").replace("*", ""))  # noqa: E731
    return {normalise_header(column.key), plain(column.en), plain(column.ar), plain(column.header_en),
            plain(column.header_ar)}


def _find_columns(rows, sheet):
    """(index of the last header row, {key: column index}) from the key row, else the Arabic or English row."""
    labels = {c.key: _labels(c) for c in sheet.columns}
    required = [c.key for c in sheet.columns if c.required]
    best, last_header = None, None
    for i, raw in enumerate(rows[:HEADER_SCAN_ROWS]):
        cells = [normalise_header(str(c).replace("[PD]", "").replace("*", "")) if c is not None else "" for c in raw]
        found = {}
        for key, names in labels.items():
            for col, text in enumerate(cells):
                if text and text in names:
                    found.setdefault(key, col)
                    break
        if len(found) * 2 >= len(sheet.columns):
            last_header = i
            exact_keys = sum(1 for c in sheet.columns if c.key in found and cells[found[c.key]] == c.key)
            score = (all(k in found for k in required), len(found), exact_keys)
            if best is None or score > best[0]:
                best = (score, found)
        elif last_header is not None:
            break
    if best is None:
        return None, {}
    return last_header, best[1]


def _sheet_by_name(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    wanted = normalise_header(name)
    for candidate in wb.sheetnames:
        if normalise_header(candidate).startswith(wanted):
            return wb[candidate]
    return None


def _example_values(sheet):
    out = []
    for example in sheet.examples:
        values = {}
        for column in sheet.columns:
            value = example.get(column.key)
            values[column.key] = (value == "yes") if column.kind == "yesno" and value is not None else (
                convert(column, value) if column.kind not in ("choice",) else value)
        out.append(values)
    return out


def _cell_problem(sheet, column, row, exc):
    detail_en, detail_ar = (("The value is not valid.", "القيمة غير صحيحة.") if column.personal else (exc.en, exc.ar))
    return problem("error", "unreadable_cell", f"{sheet.name} row {row}, {column.en}: {detail_en}",
                   f"{sheet.title_ar} الصف {row}، {column.ar}: {detail_ar}", sheet=sheet.name, row=row)


def _read_sheet(ws_rows, sheet, problems):
    data = SheetData(sheet.key, sheet.name, present=True)
    last_header, columns = _find_columns(ws_rows, sheet)
    missing = [c for c in sheet.columns if c.required and c.key not in columns]
    if last_header is None or missing:
        names = ", ".join(f"{c.en} ({c.ar})" for c in (missing or sheet.columns))
        raise ReaderError(bi(f"Sheet {sheet.name}: required columns not found: {names}.",
                             f"الورقة {sheet.title_ar}: لم يتم العثور على الأعمدة المطلوبة: {names}."))
    data.columns, data.header_row = columns, last_header + 1
    examples = _example_values(sheet)
    for i, raw in enumerate(ws_rows[last_header + 1:], start=last_header + 2):
        raw = list(raw)
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in raw):
            continue
        values, invalid = {}, set()
        for column in sheet.columns:
            col = columns.get(column.key)
            cell = raw[col] if col is not None and col < len(raw) else None
            try:
                values[column.key] = convert(column, cell)
            except CellError as exc:
                problems.append(_cell_problem(sheet, column, i, exc))
                values[column.key] = None
                invalid.add(column.key)
        record = Record(sheet.name, i, values, invalid=invalid)
        if not invalid and values in examples:
            record.example = True
            data.examples.append(record)
            problems.append(problem("error", "example_row", f"{sheet.name} row {i} is a template example: delete it.",
                                    f"{sheet.title_ar} الصف {i} مثال من القالب: احذفه.", sheet=sheet.name, row=i))
            continue
        data.records.append(record)
    return data


def read_workbook(source):
    """Parse a filled workbook (W3). Raises ReaderError for an unreadable file, a missing required sheet or column;
    everything else is a problem on ``WorkbookData.problems``."""
    from openpyxl import load_workbook

    if isinstance(source, (bytes, bytearray)):
        data, name = bytes(source), ""
    else:
        path = Path(source)
        data, name = path.read_bytes(), path.name
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for a non-xlsx file
        raise ReaderError(bi(f"Not a readable xlsx file ({type(exc).__name__}).",
                             f"الملف ليس ملف xlsx قابلًا للقراءة ({type(exc).__name__}).")) from None
    book = WorkbookData(name, hashlib.sha256(data).hexdigest(), {})
    try:
        used = set()
        for sheet in spec.SHEETS:
            ws = _sheet_by_name(wb, sheet.name)
            if ws is None:
                if sheet.required_sheet:
                    raise ReaderError(bi(f"The workbook has no {sheet.name} sheet.",
                                         f"لا يحتوي الملف على ورقة {sheet.title_ar}."))
                book.problems.append(problem("warning", "missing_sheet", f"No {sheet.name} sheet: nothing checked for it.",
                                             f"لا توجد ورقة {sheet.title_ar}: لم يُفحص شيء منها.", sheet=sheet.name))
                book.sheets[sheet.key] = SheetData(sheet.key, sheet.name)
                continue
            used.add(ws.title)
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            book.sheets[sheet.key] = _read_sheet(rows, sheet, book.problems)
        book.other_sheets = [n for n in wb.sheetnames if n not in used]
    finally:
        wb.close()
    return book


# --- adapters to the existing services --------------------------------------------------------------------------

def cutover_of(book, entity_code):
    record = book.entity(entity_code)
    if record is None:
        raise ReaderError(bi(f"Entity {entity_code} is not on the Entities sheet.",
                             f"الكيان {entity_code} غير موجود في ورقة الكيانات."))
    return record["cutover_date"]


def trial_balance_for(book, entity_code):
    """The entity's TB sheet as a parsed trial balance: closing balances (debit − credit) at its cut-over (W13)."""
    cutover_of(book, entity_code)
    rows = []
    for r in book.records("trial_balance", entity_code):
        if not r["code"]:
            continue
        closing = (r.get("debit", ZERO)) - (r.get("credit", ZERO))
        rows.append(TBRow(row=r.row, code=r["code"], name=r.get("name", ""), opening=None, movement_debit=None,
                          movement_credit=None, closing=closing))
    codes = {r.code for r in rows}
    for r in rows:
        r.is_header = any(other != r.code and other.startswith(r.code) for other in codes)
    problems = [p for p in book.problems if p.sheet == spec.SHEET_BY_KEY["trial_balance"].name]
    parsed = ParsedFile("trial_balance", f"{book.name}#{entity_code}", book.sha256, rows, problems=problems)
    parsed.has = {"code", "name", "closing_debit", "closing_credit"}
    parsed.stats["rows"] = len(rows)
    return parsed


def account_map_rows_for(book, entity_code):
    """The entity's Account mapping rows as ``create_account_map`` rows (W13)."""
    return [{"legacy_code": r["legacy_code"], "legacy_name": r.get("legacy_name", ""), "erp_code": r["erp_code"],
             "counterparty": r["counterparty"] or None, "note": r.get("note", ""), "row": r.row}
            for r in book.records("account_map", entity_code) if r["legacy_code"]]
