"""Readers for legacy exports (migration.md R1 to R3, R10): the Falcon general journal and column-mapped tables.

Readers only parse. They never touch the database, never guess an account mapping, and never drop a row silently:
an unreadable cell becomes a problem on its row. Amounts become ``Decimal`` as soon as they are read (CLAUDE.md
rule 7); a float that openpyxl returns is converted through its shortest text form.
"""
import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

ZERO = Decimal("0")
CENT = Decimal("0.01")
HEADER_SCAN_ROWS = 20

_DIACRITICS = re.compile("[ً-ٰٟۖ-ۭ]")
_SPACES = re.compile(r"\s+")
_ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬", "01234567890123456789.,")
_CODE = re.compile(r"^\d{3,5}$")


class ReaderError(ValueError):
    """The file cannot be read at all (wrong format, missing required column). Message: "English / العربية"."""


def bi(en, ar):
    return f"{en} / {ar}"


def normalise_header(text):
    """Header text for matching: trimmed, spaces collapsed, tatweel and diacritics removed, ة→ه, ى→ي, أإآ→ا."""
    if text is None:
        return ""
    text = str(text).replace("ـ", "")  # tatweel
    text = _DIACRITICS.sub("", text)
    text = text.replace("ة", "ه").replace("ى", "ي")
    for alef in "أإآٱ":
        text = text.replace(alef, "ا")
    text = text.replace("‏", "").replace("‎", "").replace("\xa0", " ")
    return _SPACES.sub(" ", text).strip().casefold()


@dataclass
class Problem:
    """One finding. ``severity`` is ``error`` (blocks loading), ``warning`` (a human looks at it) or ``info``."""

    code: str
    severity: str
    message: str  # "English / العربية"
    row: int | None = None
    entry_no: str = ""
    account_code: str = ""

    def as_dict(self):
        return {"severity": self.severity, "code": self.code, "row": self.row or "", "entry_no": self.entry_no,
                "account_code": self.account_code, "message": self.message}


def error(code, en, ar, **where):
    return Problem(code, "error", bi(en, ar), **where)


def warning(code, en, ar, **where):
    return Problem(code, "warning", bi(en, ar), **where)


# --- cell values -----------------------------------------------------------------------------------------------

class CellError(ValueError):
    def __init__(self, en, ar):
        super().__init__(bi(en, ar))
        self.en, self.ar = en, ar


def to_amount(value, stats=None):
    """A Decimal from a cell (number or text). Blank is 0. Floats go through their shortest repr (R2)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return ZERO
    if isinstance(value, bool):
        raise CellError(f"{value!r} is not an amount.", f"{value!r} ليس مبلغًا.")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        amount = Decimal(repr(value))
        rounded = amount.quantize(CENT)
        if amount != rounded and abs(amount - rounded) < Decimal("1e-9"):
            if stats is not None:
                stats["float_artifacts"] = stats.get("float_artifacts", 0) + 1
            return rounded
        return amount
    text = str(value).translate(_ARABIC_DIGITS).strip().replace(",", "").replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        amount = Decimal(text)
    except InvalidOperation:
        raise CellError(f"{value!r} is not an amount.", f"{value!r} ليس مبلغًا.") from None
    if not amount.is_finite():
        raise CellError(f"{value!r} is not an amount.", f"{value!r} ليس مبلغًا.")
    return -amount if negative else amount


def to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").translate(_ARABIC_DIGITS).strip()
    if not text:
        raise CellError("The date is empty.", "التاريخ فارغ.")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise CellError(f"{value!r} is not a date.", f"{value!r} ليس تاريخًا.")


def to_code(value):
    """An account code as text (3 to 5 digits); Excel may store it as a number."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value if value is not None else "").translate(_ARABIC_DIGITS).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not _CODE.match(text):
        raise CellError(f"Account code {value!r} is not 3 to 5 digits.", f"كود الحساب {value!r} ليس من 3 إلى 5 أرقام.")
    return text


def to_int(value, what_en, what_ar):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value if value is not None else "").translate(_ARABIC_DIGITS).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not text.isdigit():
        raise CellError(f"{what_en} {value!r} is not a whole number.", f"{what_ar} {value!r} ليس عددًا صحيحًا.")
    return int(text)


def to_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


# --- raw files -------------------------------------------------------------------------------------------------

def _read_bytes(source):
    """(bytes, name) from a path, bytes or a binary file object."""
    if isinstance(source, (bytes, bytearray)):
        return bytes(source), ""
    if isinstance(source, (str, Path)):
        path = Path(source)
        return path.read_bytes(), path.name
    data = source.read()
    return data, getattr(source, "name", "") or ""


def _xlsx_rows(data, sheet=None):
    """(rows, sheet names) of one sheet, values only, read-only (formulas are read as their cached values)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for a non-xlsx file
        raise ReaderError(bi(f"Not a readable xlsx file ({type(exc).__name__}).",
                             f"الملف ليس ملف xlsx قابلًا للقراءة ({type(exc).__name__}).")) from None
    try:
        names = wb.sheetnames
        if sheet is None:
            ws = wb[names[0]]
        elif isinstance(sheet, int):
            ws = wb[names[sheet - 1]]
        elif sheet in names:
            ws = wb[sheet]
        else:
            raise ReaderError(bi(f"Sheet {sheet!r} not found; sheets: {names}.",
                                 f"الورقة {sheet!r} غير موجودة؛ الأوراق: {names}."))
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return rows, names


def _csv_rows(data):
    """(rows, encoding). UTF-8 (with or without BOM) first, then Windows-1256 (Arabic Windows exports)."""
    for encoding in ("utf-8-sig", "cp1256"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)], encoding
    raise ReaderError(bi("The CSV file is neither UTF-8 nor Windows-1256.", "ملف CSV ليس بترميز UTF-8 ولا Windows-1256."))


def _is_xlsx(data, name):
    return data[:2] == b"PK" or name.lower().endswith((".xlsx", ".xlsm"))


def _find_header(rows, wanted, header_row=None):
    """(index of the header row, {normalised header: column index}) for the first row holding every ``wanted``."""
    candidates = [header_row - 1] if header_row else range(min(HEADER_SCAN_ROWS, len(rows)))
    for i in candidates:
        if i >= len(rows):
            break
        found = {}
        for col, cell in enumerate(rows[i]):
            key = normalise_header(cell)
            if key and key not in found:
                found[key] = col
        if all(any(normalise_header(alt) in found for alt in alts) for alts in wanted):
            return i, found
    return None, None


# --- Falcon general journal -------------------------------------------------------------------------------------

# canonical key -> (Arabic header as exported, English label, required)
FALCON_JOURNAL_COLUMNS = {
    "debit": ("مدين", "Debit", True),
    "credit": ("دائن", "Credit", True),
    "entry_no": ("رقم القيد", "Entry number", True),
    "entry_serial": ("مسلسل القيد", "Entry serial", False),
    "account_name": ("اسم الحساب", "Account name", False),
    "description": ("بيان", "Description", False),
    "journal_type": ("نوع الحركه", "Journal type", False),
    "date": ("التاريخ", "Date", True),
    "cost_centre": ("مركز التكلفة", "Cost centre", False),
    "source": ("الحركة", "Source", False),
    "source_no": ("رقم الحركة", "Source document number", False),
    "account_code": ("كود الحساب", "Account code", True),
}


@dataclass
class JournalRow:
    row: int  # 1-based row number in the sheet
    entry_no: str
    entry_serial: int | None
    date: date | None
    account_code: str
    account_name: str
    debit: Decimal
    credit: Decimal
    description: str = ""
    journal_type: str = ""
    cost_centre: str = ""
    source: str = ""
    source_no: str = ""

    def key(self):
        """Everything that identifies the line's content (duplicate detection, lines hash)."""
        return (self.entry_no, self.date, self.account_code, self.debit, self.credit, self.description,
                self.cost_centre, self.source, self.source_no, self.entry_serial)


@dataclass
class ParsedFile:
    kind: str
    name: str
    sha256: str
    rows: list
    problems: list = field(default_factory=list)
    stats: dict = field(default_factory=dict)
    columns: dict = field(default_factory=dict)  # canonical key -> header found
    extra_columns: list = field(default_factory=list)
    sheets: list = field(default_factory=list)
    encoding: str = ""
    has: set = field(default_factory=set)  # optional keys present (trial balance)


def read_falcon_journal(source):
    """Parse a Falcon general-journal export (xlsx; also CSV with the same headers). Raises ReaderError when the file
    or a required column is missing; row-level issues are ``problems``."""
    data, name = _read_bytes(source)
    sha = hashlib.sha256(data).hexdigest()
    encoding = ""
    if _is_xlsx(data, name):
        rows, sheets = _xlsx_rows(data)
    else:
        (rows, encoding), sheets = _csv_rows(data), []
    required = [(ar,) for ar, _, req in FALCON_JOURNAL_COLUMNS.values() if req]
    start, found = _find_header(rows, required)
    if start is None:
        header = rows[0] if rows else []
        present = {normalise_header(c) for c in header}
        missing = [f"{ar} ({en})" for ar, en, req in FALCON_JOURNAL_COLUMNS.values()
                   if req and normalise_header(ar) not in present]
        raise ReaderError(bi(f"Required Falcon journal columns not found: {', '.join(missing) or 'header row'}.",
                             f"لم يتم العثور على أعمدة دفتر اليومية المطلوبة: {', '.join(missing) or 'صف العناوين'}."))
    col = {}
    for key, (ar, _, _) in FALCON_JOURNAL_COLUMNS.items():
        if normalise_header(ar) in found:
            col[key] = found[normalise_header(ar)]
    known = {normalise_header(ar) for ar, _, _ in FALCON_JOURNAL_COLUMNS.values()}
    header_cells = rows[start]
    extra = [to_text(c) for c in header_cells if normalise_header(c) and normalise_header(c) not in known]
    result = ParsedFile("journal", name, sha, [], columns={k: to_text(header_cells[i]) for k, i in col.items()},
                        extra_columns=extra, sheets=sheets, encoding=encoding)
    if len(sheets) > 1:
        result.problems.append(warning("other_sheets", f"Only the first sheet was read; also present: {sheets[1:]}.",
                                       f"تمت قراءة الورقة الأولى فقط؛ توجد أيضًا: {sheets[1:]}."))
    stats = result.stats
    stats["blank_rows"] = 0
    for i, raw in enumerate(rows[start + 1:], start=start + 2):
        cells = list(raw) + [None] * (len(header_cells) - len(raw))

        def cell(key):
            return cells[col[key]] if key in col and col[key] < len(cells) else None

        if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
            stats["blank_rows"] += 1
            continue
        problems = []
        values = {}
        for key, conv in (("debit", lambda v: to_amount(v, stats)), ("credit", lambda v: to_amount(v, stats)),
                          ("date", to_date), ("account_code", to_code),
                          ("entry_no", lambda v: str(to_int(v, "Entry number", "رقم القيد")))):
            try:
                values[key] = conv(cell(key))
            except CellError as exc:
                problems.append(error("unreadable_cell", f"Row {i}, {FALCON_JOURNAL_COLUMNS[key][1]}: {exc.en}",
                                      f"الصف {i}، {FALCON_JOURNAL_COLUMNS[key][0]}: {exc.ar}", row=i))
                values[key] = ZERO if key in ("debit", "credit") else (None if key == "date" else to_text(cell(key)))
        serial = None
        if cell("entry_serial") not in (None, ""):
            try:
                serial = to_int(cell("entry_serial"), "Entry serial", "مسلسل القيد")
            except CellError as exc:
                problems.append(warning("unreadable_cell", f"Row {i}: {exc.en}", f"الصف {i}: {exc.ar}", row=i))
        if isinstance(cell("date"), datetime) and (cell("date").hour or cell("date").minute):
            stats["dates_with_time"] = stats.get("dates_with_time", 0) + 1
        for p in problems:
            p.entry_no = values["entry_no"] or ""
        result.problems += problems
        result.rows.append(JournalRow(
            row=i, entry_no=values["entry_no"] or "", entry_serial=serial, date=values["date"],
            account_code=values["account_code"] or "", account_name=to_text(cell("account_name")),
            debit=values["debit"], credit=values["credit"], description=to_text(cell("description")),
            journal_type=to_text(cell("journal_type")), cost_centre=to_text(cell("cost_centre")),
            source=to_text(cell("source")), source_no=to_text(cell("source_no")),
        ))
    stats["rows"] = len(result.rows)
    return result


# --- column-mapped tables: trial balance and chart ------------------------------------------------------------------

TB_KEYS = ("code", "name", "opening_debit", "opening_credit", "opening", "movement_debit", "movement_credit",
           "closing_debit", "closing_credit", "closing")
CHART_KEYS = ("code", "name", "name_en", "type", "parent")

# Guesses until the Falcon exports are seen (OQ-1); a column map file replaces them.
DEFAULT_TB_COLUMNS = {
    "code": ["كود الحساب", "رقم الحساب", "code", "account code"],
    "name": ["اسم الحساب", "name", "account name"],
    "opening_debit": ["رصيد أول المدة مدين", "رصيد افتتاحي مدين", "opening debit"],
    "opening_credit": ["رصيد أول المدة دائن", "رصيد افتتاحي دائن", "opening credit"],
    "movement_debit": ["حركة مدين", "الحركة مدين", "movement debit", "period debit"],
    "movement_credit": ["حركة دائن", "الحركة دائن", "movement credit", "period credit"],
    "closing_debit": ["الرصيد الختامي مدين", "رصيد آخر المدة مدين", "رصيد مدين", "closing debit"],
    "closing_credit": ["الرصيد الختامي دائن", "رصيد آخر المدة دائن", "رصيد دائن", "closing credit"],
}
DEFAULT_CHART_COLUMNS = {
    "code": ["كود الحساب", "رقم الحساب", "code"],
    "name": ["اسم الحساب", "name"],
    "name_en": ["الاسم بالانجليزي", "name_en", "english name"],
    "type": ["نوع الحساب", "type"],
    "parent": ["الحساب الرئيسي", "parent"],
}


@dataclass
class TBRow:
    row: int
    code: str
    name: str
    opening: Decimal | None  # signed: debit positive
    movement_debit: Decimal | None
    movement_credit: Decimal | None
    closing: Decimal
    is_header: bool = False

    @property
    def movement(self):
        if self.movement_debit is None and self.movement_credit is None:
            return None
        return (self.movement_debit or ZERO) - (self.movement_credit or ZERO)


@dataclass
class ChartRow:
    row: int
    code: str
    name: str
    name_en: str = ""
    type: str = ""
    parent: str = ""


def load_column_map(source):
    """A column map from a JSON file, a JSON text or a dict: ``{"columns": {...}, "sheet": ..., "header_row": n}``
    (or just the ``columns`` dict)."""
    if source is None:
        return None
    if isinstance(source, dict):
        spec = source
    else:
        text = Path(source).read_text(encoding="utf-8") if not str(source).lstrip().startswith("{") else str(source)
        try:
            spec = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ReaderError(bi(f"The column map is not valid JSON ({exc.msg}).",
                                 f"خريطة الأعمدة ليست JSON صالحًا ({exc.msg}).")) from None
    if "columns" not in spec:
        spec = {"columns": spec}
    return spec


def _resolve_columns(rows, columns, allowed, required, header_row):
    """(header row index, {key: column index}) for a column map; ints are 1-based column numbers."""
    unknown = set(columns) - set(allowed)
    if unknown:
        raise ReaderError(bi(f"Unknown column map keys: {sorted(unknown)}; allowed: {list(allowed)}.",
                             f"مفاتيح غير معروفة في خريطة الأعمدة: {sorted(unknown)}؛ المسموح: {list(allowed)}."))
    alts = {k: (v if isinstance(v, list) else [v]) for k, v in columns.items()}
    named = [[a for a in alts[k] if not isinstance(a, int)] for k in required if k in alts]
    named = [n for n in named if n]
    if named:
        start, found = _find_header(rows, named, header_row)
    else:
        start, found = (header_row - 1 if header_row else 0), {}
    if start is None:
        raise ReaderError(bi(f"Header row not found; looked for: {named}.", f"لم يتم العثور على صف العناوين؛ المطلوب: {named}."))
    if not found:
        found = {normalise_header(c): i for i, c in enumerate(rows[start] if start < len(rows) else [])}
    resolved = {}
    for key, options in alts.items():
        for alt in options:
            if isinstance(alt, int):
                resolved[key] = alt - 1
                break
            if normalise_header(alt) in found:
                resolved[key] = found[normalise_header(alt)]
                break
    missing = [k for k in required if k not in resolved]
    if missing:
        raise ReaderError(bi(f"Required columns not found: {missing}.", f"لم يتم العثور على الأعمدة المطلوبة: {missing}."))
    return start, resolved


def _table(source, spec):
    data, name = _read_bytes(source)
    sha = hashlib.sha256(data).hexdigest()
    if _is_xlsx(data, name):
        rows, sheets = _xlsx_rows(data, spec.get("sheet"))
        encoding = ""
    else:
        (rows, encoding), sheets = _csv_rows(data), []
    return data, name, sha, rows, sheets, encoding


def read_trial_balance(source, column_map=None):
    """Parse a trial balance with a column map (R3, R10). Rows whose code is a prefix of another row's code are
    marked as headers (``is_header``); the validation excludes them and checks their subtotals."""
    spec = load_column_map(column_map) or {"columns": DEFAULT_TB_COLUMNS}
    _, name, sha, rows, sheets, encoding = _table(source, spec)
    columns = spec["columns"]
    if "closing" not in columns and not {"closing_debit", "closing_credit"} <= set(columns):
        raise ReaderError(bi("The trial-balance column map needs closing (signed) or closing_debit and closing_credit.",
                             "خريطة أعمدة ميزان المراجعة تحتاج إلى الرصيد الختامي (بإشارة) أو الختامي المدين والدائن."))
    start, col = _resolve_columns(rows, columns, TB_KEYS, ["code"], spec.get("header_row"))
    if "closing" not in col and not {"closing_debit", "closing_credit"} <= set(col):
        raise ReaderError(bi("No closing balance column found in the trial balance.",
                             "لم يتم العثور على عمود الرصيد الختامي في ميزان المراجعة."))
    result = ParsedFile("trial_balance", name, sha, [], columns={k: to_text(rows[start][i]) if start < len(rows) and
                        i < len(rows[start]) else f"#{i + 1}" for k, i in col.items()}, sheets=sheets, encoding=encoding)
    result.has = set(col)
    stats = result.stats
    for i, raw in enumerate(rows[start + 1:], start=start + 2):
        def cell(key):
            return raw[col[key]] if key in col and col[key] < len(raw) else None

        if all(c is None or (isinstance(c, str) and not c.strip()) for c in raw):
            continue
        try:
            code = to_code(cell("code"))
        except CellError as exc:
            if not to_text(cell("code")):  # a totals or title row without a code: reported, not loaded
                stats["rows_without_code"] = stats.get("rows_without_code", 0) + 1
                continue
            result.problems.append(error("unreadable_cell", f"Row {i}: {exc.en}", f"الصف {i}: {exc.ar}", row=i))
            continue

        def amount(key):
            if key not in col:
                return None
            try:
                return to_amount(cell(key), stats)
            except CellError as exc:
                result.problems.append(error("unreadable_cell", f"Row {i}, {key}: {exc.en}", f"الصف {i}، {key}: {exc.ar}",
                                             row=i, account_code=code))
                return ZERO

        def signed(single, dr, cr):
            if single in col:
                return amount(single)
            if dr in col or cr in col:
                return (amount(dr) or ZERO) - (amount(cr) or ZERO)
            return None

        result.rows.append(TBRow(
            row=i, code=code, name=to_text(cell("name")),
            opening=signed("opening", "opening_debit", "opening_credit"),
            movement_debit=amount("movement_debit"), movement_credit=amount("movement_credit"),
            closing=signed("closing", "closing_debit", "closing_credit") or ZERO,
        ))
    codes = {r.code for r in result.rows}
    for r in result.rows:
        r.is_header = any(other != r.code and other.startswith(r.code) for other in codes)
    stats["rows"] = len(result.rows)
    return result


def read_chart(source, column_map=None):
    """Parse a chart-of-accounts export with a column map (R3)."""
    spec = load_column_map(column_map) or {"columns": DEFAULT_CHART_COLUMNS}
    _, name, sha, rows, sheets, encoding = _table(source, spec)
    start, col = _resolve_columns(rows, spec["columns"], CHART_KEYS, ["code", "name"], spec.get("header_row"))
    result = ParsedFile("chart", name, sha, [], sheets=sheets, encoding=encoding)
    for i, raw in enumerate(rows[start + 1:], start=start + 2):
        def cell(key):
            return raw[col[key]] if key in col and col[key] < len(raw) else None

        if all(c is None or (isinstance(c, str) and not c.strip()) for c in raw):
            continue
        try:
            code = to_code(cell("code"))
        except CellError as exc:
            result.problems.append(error("unreadable_cell", f"Row {i}: {exc.en}", f"الصف {i}: {exc.ar}", row=i))
            continue
        result.rows.append(ChartRow(i, code, to_text(cell("name")), to_text(cell("name_en")), to_text(cell("type")),
                                    to_text(cell("parent"))))
    result.stats["rows"] = len(result.rows)
    return result


def read_account_map_csv(source):
    """Rows of a ``docs/migration/<entity>-account-map.csv`` file: legacy_code, legacy_name, erp_code, counterparty,
    note (header names in English; UTF-8)."""
    data, _ = _read_bytes(source)
    rows, _ = _csv_rows(data)
    if not rows:
        raise ReaderError(bi("The account map file is empty.", "ملف خريطة الحسابات فارغ."))
    header = [normalise_header(h) for h in rows[0]]
    for needed in ("legacy_code", "erp_code"):
        if needed not in header:
            raise ReaderError(bi(f"The account map needs a {needed} column.", f"خريطة الحسابات تحتاج إلى العمود {needed}."))
    out = []
    for i, raw in enumerate(rows[1:], start=2):
        record = {header[j]: (raw[j].strip() if j < len(raw) else "") for j in range(len(header))}
        if not any(record.values()):
            continue
        record["row"] = i
        out.append(record)
    return out
