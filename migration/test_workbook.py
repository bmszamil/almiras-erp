"""migration.md §11 W1, W3 (generator and reader; AT-16 to AT-18). Synthetic workbooks only, built in the test."""
import io
import re
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from migration import readers, workbook
from migration import testing as t
from migration import workbook_spec as spec

D = Decimal
ARABIC = re.compile("[؀-ۿ]")


def template():
    return load_workbook(io.BytesIO(workbook.build_workbook()))


# --- generator (AT-16) ---------------------------------------------------------------------------------------------

def test_at16_sheets_in_order_with_bilingual_headers_keys_and_markers():
    wb = template()
    assert wb.sheetnames == ["Instructions", *[s.name for s in spec.SHEETS], "Checklist", "ERP chart", "Lists"]
    for sheet in spec.SHEETS:
        ws = wb[sheet.name]
        assert sheet.title_ar in ws["A1"].value and ARABIC.search(ws["A2"].value)
        for i, column in enumerate(sheet.columns, start=1):
            en, ar, key = (ws.cell(row=r, column=i).value for r in (spec.EN_ROW, spec.AR_ROW, spec.KEY_ROW))
            assert (en, ar, key) == (column.header_en, column.header_ar, column.key)
            assert ARABIC.search(ar)
            assert en.endswith("*") or "* " in en if column.required else "*" not in en
            assert ("[PD]" in en and "[PD]" in ar) == column.personal
            fill = ws.cell(row=spec.EN_ROW, column=i).fill.fgColor.rgb[-6:]
            assert fill == (workbook.FILL_PERSONAL if column.personal else workbook.FILL_REQUIRED if column.required
                            else workbook.FILL_HEADER)
    employees = spec.SHEET_BY_KEY["employees"]
    assert {c.key for c in employees.columns if c.personal} >= {"full_name_ar", "full_name_en", "id_number",
                                                                "date_of_birth", "gender", "iban", "basic_salary"}
    assert not employees.column("employee_no").personal and not employees.column("hire_date").personal
    instructions = wb["Instructions"]
    assert ARABIC.search(instructions["B3"].value) and "Falcon" in instructions["A3"].value


def test_at16_dropdowns_number_formats_and_lists():
    wb = template()
    names = wb.defined_names
    assert names["L_entity"].attr_text == "'Entities'!$A$6:$A$2005"
    assert names["L_yes_no"].attr_text.startswith("'Lists'!")
    lists = wb["Lists"]
    headers = [c.value for c in lists[1]]
    yes_no = [lists.cell(row=r, column=headers.index("yes_no") + 1).value for r in (2, 3)]
    assert yes_no == ["Yes / نعم", "No / لا"]
    currencies = [lists.cell(row=r, column=headers.index("currency") + 1).value for r in range(2, 40)]
    assert {"SAR", "USD", "KWD", "JPY"} <= set(currencies)
    erp = [lists.cell(row=r, column=headers.index("erp_code") + 1).value for r in range(2, 120)]
    assert "1200" in erp and "1000" not in erp and "3400" not in erp  # headers and report-only left out

    ws = wb["Open receivables"]
    sheet = spec.SHEET_BY_KEY["open_receivables"]
    letter = {c.key: chr(ord("A") + i) for i, c in enumerate(sheet.columns)}
    by_range = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    assert by_range[f"{letter['entity']}6:{letter['entity']}2005"].formula1 == "L_entity"
    assert by_range[f"{letter['document_type']}6:{letter['document_type']}2005"].formula1 == "L_ar_document"
    assert by_range[f"{letter['currency']}6:{letter['currency']}2005"].formula1 == "L_currency"
    assert by_range[f"{letter['sar_amount']}6:{letter['sar_amount']}2005"].type == "decimal"
    assert by_range[f"{letter['document_date']}6:{letter['document_date']}2005"].type == "date"
    assert ws.column_dimensions[letter["sar_amount"]].number_format == "#,##0.00"
    assert ws.column_dimensions[letter["amount"]].number_format == "#,##0.000"
    assert ws.column_dimensions[letter["document_date"]].number_format == "yyyy-mm-dd"
    assert ws.column_dimensions[letter["document_no"]].number_format == "@"  # identifiers stay text
    entities = wb["Entities"]
    yes = {str(dv.sqref): dv.formula1 for dv in entities.data_validations.dataValidation}
    assert "L_yes_no" in yes.values() and "L_legacy_system" in yes.values()
    checklist = wb["Checklist"]
    assert ARABIC.search(checklist.cell(row=spec.FIRST_DATA_ROW, column=3).value)
    assert "المحاسب" in str(checklist.cell(row=spec.FIRST_DATA_ROW + 1, column=4).value)


def test_at16_examples_are_fake_and_flagged_when_left_in():
    for sheet in spec.SHEETS:
        assert 2 <= len(sheet.examples) <= 3, sheet.key
        for example in sheet.examples:
            assert set(example) == {c.key for c in sheet.columns}, sheet.key
    for example in spec.SHEET_BY_KEY["entities"].examples:
        assert "(fake)" in example["name_en"] and "وهمية" in example["name_ar"]
    for iban in (spec.EXAMPLE_IBAN_1, spec.EXAMPLE_IBAN_2, spec.EXAMPLE_IBAN_3):
        from banking import iban as rules
        assert rules.problems(iban, country="SA") == []
    book = workbook.read_workbook(workbook.build_workbook())
    flagged = [(p.sheet, p.row) for p in book.problems if p.code == "example_row"]
    assert len(flagged) == sum(len(s.examples) for s in spec.SHEETS)
    assert all(not book.sheets[s.key].records for s in spec.SHEETS)
    assert [p.code for p in book.problems if p.code != "example_row"] == []


def test_at16_committed_template_is_the_generated_one():
    """Regenerate with ``manage.py migration_template`` after changing workbook_spec.py."""
    committed = load_workbook(workbook.DEFAULT_PATH)
    fresh = template()
    assert committed.sheetnames == fresh.sheetnames
    for name in fresh.sheetnames:
        a, b = committed[name], fresh[name]
        assert [[c.value for c in row] for row in a.iter_rows()] == [[c.value for c in row] for row in b.iter_rows()], name
        assert sorted((str(v.sqref), v.formula1) for v in a.data_validations.dataValidation) == \
            sorted((str(v.sqref), v.formula1) for v in b.data_validations.dataValidation), name
    assert {n: d.attr_text for n, d in committed.defined_names.items()} == \
        {n: d.attr_text for n, d in fresh.defined_names.items()}


def test_choice_codes_follow_the_erp_models():
    from banking.models import BankAccount, StatementFormat
    from hr.models import Employee, EmploymentBasis
    from invoicing.models import Customer
    from invoicing.reference_seed import WHT_CATEGORIES_V1
    from ledger.models import Account

    codes = lambda key: [c[0] for c in spec.CHOICES[key]]  # noqa: E731
    assert codes("party_kind") == list(Customer.Kind.values)
    assert codes("account_type") == list(Account.Type.values)
    assert codes("id_type") == list(Employee.IdType.values)
    assert codes("gender") == list(Employee.Gender.values)
    assert codes("gosi_cohort") == list(Employee.GosiCohort.values)
    assert codes("employment_basis") == list(EmploymentBasis.values)
    assert codes("sharia_basis") == list(BankAccount.ShariaBasis.values)
    assert codes("statement_format") == list(StatementFormat.values)
    assert codes("wht_category") == [c[0] for c in WHT_CATEGORIES_V1]
    assert codes("bank")[:-1] == list(spec.BANKS_V1)


# --- reader (AT-17, AT-18) -----------------------------------------------------------------------------------------

def test_at17_filled_workbook_reads_typed_values_and_keeps_arabic():
    book = workbook.read_workbook(t.filled_workbook())
    assert book.entity_codes() == ["ALM", "SUB"] and book.entity("SUB")["parent_entity"] == "ALM"
    assert book.entity("ALM")["cutover_date"] == date(2026, 12, 31) and book.entity("ALM")["sharia_enforced"] is True
    tb = book.records("trial_balance", "ALM")
    assert (tb[0]["code"], tb[0]["name"], tb[0]["debit"], tb[0]["credit"]) == ("1101", "بنك تجريبي جاري", D("50000.00"), None)
    customers = {r["code"]: r for r in book.records("customers", "ALM")}
    assert customers["C002"]["kind"] == "individual" and customers["C002"]["name_ar"] == "عميل اختبار فرد"
    assert customers["C-SUB"]["group_entity"] == "SUB"
    item = book.records("open_receivables", "ALM")[2]
    assert (item["currency"], item["amount"], item["sar_amount"]) == ("USD", D("2400.00"), D("9000.00"))
    assert book.records("bank_accounts", "ALM")[0]["iban"].startswith("SA")
    assert book.records("employees", "ALM")[0]["date_of_birth"] == date(1990, 1, 1)
    assert [p.code for p in book.problems] == [] and len(book.sha256) == 64


def test_at17_choices_as_codes_english_or_arabic_and_codes_typed_as_numbers():
    rows = t.workbook_rows()
    rows["falcon_chart"][1]["type"] = "Asset"
    rows["falcon_chart"][2]["type"] = "أصل"
    rows["falcon_chart"][3]["postable"] = "نعم"
    for r in rows["trial_balance"]:
        r["code"] = int(r["code"])  # Excel stores a typed code as a number
    book = workbook.read_workbook(t.filled_workbook(rows, labels=False))
    chart = book.records("falcon_chart", "ALM")
    assert [r["type"] for r in chart[:4]] == ["asset"] * 4 and chart[3]["postable"] is True
    assert book.records("trial_balance", "ALM")[0]["code"] == "1101"
    assert book.entity("ALM")["legacy_system"] == "falcon"


def test_at17_reads_without_the_key_row():
    book = workbook.read_workbook(t.filled_workbook(drop_key_row=True))
    assert len(book.records("trial_balance")) == 15 and not book.problems


def test_at18_missing_required_column_or_sheet_is_refused_in_both_languages():
    with pytest.raises(readers.ReaderError) as exc:
        workbook.read_workbook(t.filled_workbook(drop_columns=[("trial_balance", "code")]))
    assert "Falcon account code" in str(exc.value) and "كود الحساب في فالكون" in str(exc.value)
    with pytest.raises(readers.ReaderError) as exc:
        workbook.read_workbook(t.filled_workbook(drop_sheets=["account_map"]))
    assert "Account mapping" in str(exc.value) and "ربط الحسابات" in str(exc.value)
    book = workbook.read_workbook(t.filled_workbook(drop_sheets=["employees"]))
    assert [(p.code, p.severity) for p in book.problems] == [("missing_sheet", "warning")]
    with pytest.raises(readers.ReaderError):
        workbook.read_workbook(b"PK\x03\x04 broken")


def test_at18_example_rows_left_in_are_errors_and_not_data():
    book = workbook.read_workbook(t.filled_workbook(keep_examples=True))
    assert {p.code for p in book.problems} == {"example_row"}
    assert book.entity_codes() == ["ALM", "SUB"]  # EXA / EXB examples are not data


def test_unreadable_cells_are_row_problems_and_personal_values_are_not_repeated():
    rows = t.workbook_rows()
    rows["trial_balance"][0]["debit"] = D("50000.005")
    rows["employees"][0]["id_number"] = "SECRET-ID-" + "9" * 20  # longer than 20 characters
    rows["employees"][0]["date_of_birth"] = "secret-birthday"
    rows["customers"][0]["vat_number"] = "300000000000000"
    book = workbook.read_workbook(t.filled_workbook(rows))
    found = {(p.sheet, p.row) for p in book.problems if p.code == "unreadable_cell"}
    assert found == {("Trial balance", 6), ("Employees", 6), ("Customers", 6)}
    text = " ".join(p.message for p in book.problems)
    assert "50000.005" in text and "SECRET" not in text and "secret-birthday" not in text
    assert book.records("employees", "ALM")[0]["id_number"] is None  # kept as a row, the cell unreadable


def test_adapters_to_the_existing_loaders():
    book = workbook.read_workbook(t.filled_workbook())
    tb = workbook.trial_balance_for(book, "ALM")
    assert tb.kind == "trial_balance" and {"closing_debit", "closing_credit"} <= tb.has
    rows = {r.code: r for r in tb.rows}
    assert rows["1101"].closing == D("50000.00") and rows["2601"].closing == D("-100000.00")
    assert rows["1101"].opening is None and not any(r.is_header for r in tb.rows)
    assert sum((r.closing for r in tb.rows), D(0)) == 0
    mapping = {r["legacy_code"]: r for r in workbook.account_map_rows_for(book, "ALM")}
    assert mapping["1202"]["erp_code"] == "1210" and mapping["1202"]["counterparty"] == "SUB"
    assert mapping["1101"]["counterparty"] is None and mapping["1101"]["legacy_name"] == "بنك تجريبي جاري"
    with pytest.raises(readers.ReaderError):
        workbook.trial_balance_for(book, "XYZ")
