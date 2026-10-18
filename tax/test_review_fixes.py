"""Regression tests for the code review and compliance findings on the VAT slice (tax-returns.md §10.4).

Each test names the finding it covers ("item n" = code review, "Cn" = compliance check). The reviewer's probes
(scratch ``test_review_probes.py`` / ``test_review_pg_probes.py``) are ported here with the corrected outcome.
"""
from datetime import date

import pytest

from core import access
from ledger import services as ledger
from ledger.models import JournalEntry
from tax import services
from tax.conftest import D, approved, bill, close_q1, filed, sale
from tax.models import VatReturn

pytestmark = pytest.mark.django_db
Q1, Q2 = date(2027, 1, 1), date(2027, 4, 1)


# --- item 5: registered tax documents are reversed only through tax.services ---------------------------------

def _paid_q1(t):
    sale(t)
    close_q1(t)
    ret = filed(t, approved(t))
    pay = services.create_vat_payment(ret, user=t.a1, amount_sar="1500.00", payment_date=date(2027, 4, 25),
                                      bank_account="1100")
    services.submit_tax_document(pay, user=t.a1)
    return ret, services.approve_tax_document(pay, user=t.c1)


def test_item5_generic_ledger_reversal_of_a_tax_payment_is_refused(t):
    ret, pay = _paid_q1(t)
    entry = pay.journal_entry
    with pytest.raises(access.AccessDenied) as exc:
        ledger.reverse_entry(entry, access.require(t.a1, t.alm, "ledger.reverse_entry", document=entry), reason="x")
    assert exc.value.code == "document_reversal_only"
    with pytest.raises(access.AccessDenied) as exc:  # its own permission, but outside reverse_tax_document
        ledger.reverse_entry(entry, access.require(t.a1, t.alm, "tax.reverse_document", document=pay), reason="x")
    assert exc.value.code == "reverse_through_owner"
    assert JournalEntry.objects.get(pk=entry.pk).status == "posted"
    assert VatReturn.objects.get(pk=ret.pk).status == "settled"
    services.reverse_tax_document(pay, user=t.a1, reason="Paid from the wrong account")
    assert VatReturn.objects.get(pk=ret.pk).status == "filed" and services.open_amount(ret, user=t.a1) == D("1500.00")


def test_item5_t1_offset_is_never_reversed(t):
    sale(t)
    bill(t, [("5300", "1", "2000.00")])
    close_q1(t)
    ret = approved(t)
    entry = ret.settlement_entry
    with pytest.raises(access.AccessDenied) as exc:
        ledger.reverse_entry(entry, access.require(t.a1, t.alm, "ledger.reverse_entry", document=entry), reason="x")
    assert exc.value.code == "document_reversal_only"
    with pytest.raises(access.AccessDenied) as exc:
        ledger.reverse_entry(entry, access.require(t.a1, t.alm, "tax.reverse_document", document=ret), reason="x")
    assert exc.value.code == "grant_wrong_permission"
    assert JournalEntry.objects.get(pk=entry.pk).status == "posted"


def test_item5_invoicing_documents_still_reverse_through_invoicing(t):
    from invoicing import services as invoicing

    doc = bill(t, [("5300", "1", "1000.00")], on=date(2027, 3, 10))
    with pytest.raises(access.AccessDenied) as exc:
        ledger.reverse_entry(doc.journal_entry, access.require(t.a1, t.alm, "ledger.reverse_entry",
                                                               document=doc.journal_entry))
    assert exc.value.code == "document_reversal_only"
    done = invoicing.reverse_document(doc, user=t.a1, date=date(2027, 3, 20), reason="wrong vendor")
    assert done.status == "reversed"


# --- item 1: the Art. 51 true-up when a quarter crosses 31 December -------------------------------------------

def test_item1_last_return_of_year_finds_31_december_inside_the_period():
    assert services._period_end(date(2027, 11, 1), "quarterly", 2) == date(2028, 1, 31)
    assert services._last_return_of_year(VatReturn(period_start=date(2027, 11, 1), period_end=date(2028, 1, 31))) == 2027
    assert services._last_return_of_year(VatReturn(period_start=date(2027, 10, 1), period_end=date(2027, 12, 31))) == 2027
    assert services._last_return_of_year(VatReturn(period_start=date(2028, 2, 1), period_end=date(2028, 4, 30))) is None


class YearArt51:
    """Art. 51 port stub that records the years asked for and reports provisional recovery in 2027."""
    years = []

    def year_status(self, entity, year):
        from tax import art51_port

        self.years.append(year)
        return art51_port.Art51Status(applies=year == 2027)


def test_item1_quarter_start_month_2_blocks_approval_of_the_nov_jan_return(t, settings):
    from core.access.testing import member
    from invoicing.conftest import make_entity
    from ledger.services_setup import create_monthly_periods
    from tax.conftest import close, reviewed

    settings.TAX_ART51_PORT = "tax.test_review_fixes.YearArt51"
    new = make_entity("QSM", "full", t.alm, vat="344444444400003", cr="1010000011",
                      vat_effective_date=date(2027, 11, 1))
    create_monthly_periods(new, 2028)
    for name, role in (("a1", "accountant"), ("a2", "accountant"), ("c1", "approver")):
        member(getattr(t, name), new, role)
    services.change_tax_settings(new, user=t.c1, vat_quarter_start_month=2)
    t.today.set(date(2028, 2, 10))
    ret = services.compute_draft(new, date(2027, 11, 1), user=t.a1)
    assert (ret.period_start, ret.period_end) == (date(2027, 11, 1), date(2028, 1, 31))
    ret = reviewed(t, services.submit_for_review(ret, user=t.a1))
    close(t, "2027-11", "2027-12", "2028-01", entity=new)
    YearArt51.years = []
    with pytest.raises(services.TaxError) as exc:
        services.approve_return(ret, user=t.c1)
    assert exc.value.code == "art51_true_up_missing" and "2027" in str(exc.value)
    assert 2027 in YearArt51.years and 2028 not in YearArt51.years


# --- item 2: approval re-runs the exceptions and the tie-out --------------------------------------------------

def test_item2_journal_posted_after_review_blocks_approval(t):
    from tax.conftest import journal, prepared, reviewed

    sale(t)
    ret = reviewed(t, prepared(t))
    journal(t, date(2027, 3, 20), [{"account": "1100", "debit": "5000.00"}, {"account": "2200", "credit": "5000.00"}])
    close_q1(t)
    with pytest.raises(services.TaxError) as exc:
        services.approve_return(ret, user=t.c1)
    assert exc.value.code == "exceptions_open" and "untagged_vat_journal" in str(exc.value)
    assert "tie_out_difference" in str(exc.value) and VatReturn.objects.get(pk=ret.pk).status == "reviewed"


def test_item2_source_hash_covers_the_vat_account_movements(t):
    sale(t)
    close_q1(t)
    ret = approved(t)
    boxes_only = __import__("hashlib").sha256(services._canonical_of_stored(ret).encode()).hexdigest()
    assert ret.source_hash == services._source_hash(ret) != boxes_only


# --- item 3: legacy imports only before the go-live date, with two other people -------------------------------

def test_item3_one_accountant_cannot_file_an_import_alone(t):
    from tax.conftest import imported

    with pytest.raises(services.TaxError) as exc:  # the probe: a 2027 (ERP) period with sales, imported as filed
        sale(t)
        services.import_filed_return(t.alm, user=t.a1, period_start=Q1, zatca_filing_reference="FAKE",
                                     filed_on=date(2027, 4, 2), boxes={13: {"vat": "-90000.00"}, 16: {"vat": "-90000.00"}})
    assert exc.value.code == "import_after_golive"
    draft = services.import_filed_return(t.alm, user=t.a1, period_start=date(2026, 10, 1), zatca_filing_reference="F",
                                         filed_on=date(2027, 1, 20), boxes={13: {"vat": "-4000.00"},
                                                                            16: {"vat": "-4000.00"}})
    assert (draft.status, draft.source, draft.schedule_id) == ("draft", "import", None)
    assert services.compute_draft(t.alm, Q1, user=t.a1).boxes.get(box_no=15).vat_sar == 0  # not filed yet
    with pytest.raises(services.TaxError) as exc:
        services.recompute_draft(draft, user=t.a1)
    assert exc.value.code == "import_not_recomputed"
    draft = services.submit_for_review(draft, user=t.a1)
    with pytest.raises(access.AccessDenied) as denied:
        services.complete_review(draft, user=t.a1)
    assert denied.value.code == "self_approval"
    draft = services.complete_review(draft, user=t.a2)
    with pytest.raises(access.AccessDenied) as denied:
        services.approve_return(draft, user=t.a2)
    filed_ = services.approve_return(draft, user=t.c1)
    assert (filed_.status, filed_.approved_by, filed_.reviewed_by) == ("filed", t.c1, t.a2)
    assert services.recompute_draft(services.compute_draft(t.alm, Q1, user=t.a1), user=t.a1).boxes.get(
        box_no=15).vat_sar == D("4000.00")
    with pytest.raises(services.TaxError) as exc:
        imported(t, date(2026, 10, 1))
    assert exc.value.code == "no_such_period"


def test_item3_import_needs_a_golive_date_and_a_period_without_erp_sources(t):
    from core.access.testing import member
    from invoicing.conftest import make_entity

    new = make_entity("LEG", "full", t.alm, vat="355555555500003", cr="1010000012")
    member(t.a1, new, "accountant")
    with pytest.raises(services.TaxError) as exc:
        services.import_filed_return(new, user=t.a1, period_start=date(2026, 10, 1), zatca_filing_reference="F",
                                     filed_on=date(2027, 1, 20), boxes={})
    assert exc.value.code == "legacy_cutoff_missing"
    sale(t, on=date(2026, 11, 10))
    with pytest.raises(services.TaxError) as exc:
        services.import_filed_return(t.alm, user=t.a1, period_start=date(2026, 10, 1), zatca_filing_reference="F",
                                     filed_on=date(2027, 1, 20), boxes={})
    assert exc.value.code == "import_period_has_sources"


def test_item3_g4_golive_must_start_a_period_and_nothing_is_scheduled_before_it(t):
    with pytest.raises(services.TaxError) as exc:
        services.change_tax_settings(t.alm, user=t.c1, vat_golive_date=date(2027, 2, 1))
    assert exc.value.code == "golive_not_period_start"
    rows = services.generate_schedule(t.alm, user=t.a1, until=date(2027, 6, 30))
    assert rows[0].period_start == Q1 and len(rows) == 2


# --- item 4: correcting entries must match box 14 ---------------------------------------------------------------

def _filed_q4_2026(t):
    from tax.conftest import imported

    return imported(t, date(2026, 10, 1), ref="OLD", filed_on=date(2027, 1, 20))


def test_item4_a_correction_cannot_hide_an_unrelated_invoice(t):
    from tax.conftest import box, confirm_tax_points, prepared

    old = _filed_q4_2026(t)
    big = sale(t, [("rent_commercial", "1", "100000.00")])
    services.record_correction(old, user=t.a1, discovered_on=date(2027, 2, 1), description_en="x", description_ar="س",
                               tax_difference_sar="1.00", correcting_entries=[big.journal_entry])
    ret = prepared(t)
    assert box(ret, 1) == (D("0"), D("0"), D("0")) and box(ret, 14)[2] == D("1.00")
    found = {e["kind"]: e for e in services.exceptions(ret, user=t.a2)}
    assert found["correcting_entry"]["reference"] == f"je:{big.journal_entry_id}"
    assert found["correction_mismatch"]["difference"] == D("1.00") - D("15000.00")
    confirm_tax_points(t, ret)
    with pytest.raises(services.TaxError) as exc:
        services.complete_review(ret, user=t.a2)
    assert exc.value.code == "exceptions_open"


def test_item4_a_correction_without_a_correcting_entry_blocks_until_one_is_linked(t):
    from tax.conftest import confirm_tax_points, journal, prepared

    old = _filed_q4_2026(t)
    corr = services.record_correction(old, user=t.a1, discovered_on=date(2027, 2, 1), description_en="Missed VAT",
                                      description_ar="ضريبة لم تُقرّ", tax_difference_sar="100.00")
    ret = prepared(t)
    kinds = {e["kind"] for e in services.exceptions(ret, user=t.a2)}
    assert "correction_without_entry" in kinds
    with pytest.raises(services.TaxError) as exc:  # it is never acknowledged: 2200 would drift when Q1 is paid
        services.acknowledge_exception(ret, user=t.a2, kind="correction_without_entry", reference=f"corr:{corr.pk}",
                                       decision="difference_explained", reason_ar="سبب")
    assert exc.value.code == "no_such_exception"
    with pytest.raises(services.TaxError):
        services.complete_review(ret, user=t.a2)
    services.return_to_draft(ret, user=t.a2, reason_ar="ربط قيد التصحيح")
    entry = journal(t, date(2027, 2, 2), [{"account": "1100", "debit": "100.00"}, {"account": "2200", "credit": "100.00"}])
    services.link_correcting_entries(corr, user=t.a1, entries=[entry])
    ret = services.submit_for_review(services.recompute_draft(ret, user=t.a1), user=t.a1)
    services.acknowledge_exception(ret, user=t.a2, kind="correcting_entry", reference=f"je:{entry.pk}",
                                   decision="confirmed_correction", reason_ar="قيد التصحيح")
    confirm_tax_points(t, ret)
    assert services.complete_review(ret, user=t.a2).status == "reviewed"
    tie = services.tie_out(ret, user=t.a2)
    assert (tie["difference_output"], tie["difference_input"]) == (0, 0)


# --- item 6: dated frequency history; the schedule does not depend on when it is generated --------------------

def _new_quarterly(t, code="FRQ", vat="366666666600003", cr="1010000013"):
    from core.access.testing import make_user, member
    from invoicing.conftest import make_entity

    new = make_entity(code, "full", t.alm, vat=vat, cr=cr, vat_effective_date=date(2027, 1, 1))
    acc = make_user(f"{code.lower()}_acc")
    member(acc, new, "accountant")
    return new, acc


def test_item6_a_late_generation_keeps_the_quarter_that_started_before_the_change(t):
    from core import services as core_services

    new, acc = _new_quarterly(t)
    t.today.set(date(2027, 3, 1))
    assert [(r.period_start, r.period_end) for r in services.generate_schedule(new, user=acc, until=date(2027, 3, 31))] \
        == [(Q1, date(2027, 3, 31))]
    t.today.set(date(2027, 5, 15))
    core_services.change_entity_settings(new, user=t.mp, reason="Monthly from now", vat_return_period="monthly")
    t.today.set(date(2027, 8, 10))  # generated long after the change: Q2 started before it, so it stays a quarter
    rows = services.generate_schedule(new, user=acc, until=date(2027, 8, 31))
    assert [(r.period_start, r.period_end, r.frequency) for r in rows][1:] == [
        (Q2, date(2027, 6, 30), "quarterly"), (date(2027, 7, 1), date(2027, 7, 31), "monthly"),
        (date(2027, 8, 1), date(2027, 8, 31), "monthly")]
    change = new.vat_frequency_changes.get()
    assert (change.previous_frequency, change.frequency, change.changed_on, change.recorded_by) == (
        "quarterly", "monthly", date(2027, 5, 15), t.mp)
    change._service_write = True
    with pytest.raises(services.TaxError):  # append-only
        change.save()


def test_item6_change_entity_settings_regenerates_a_pregenerated_quarter(t):
    from core import services as core_services

    new, acc = _new_quarterly(t)
    t.today.set(date(2027, 3, 1))
    services.generate_schedule(new, user=acc, until=date(2027, 9, 30))
    t.today.set(date(2027, 6, 20))
    core_services.change_entity_settings(new, user=t.mp, reason="Monthly from July", vat_return_period="monthly")
    rows = services.schedule(new, user=acc)  # regenerated by the change itself, before July starts
    assert [(r.period_start, r.period_end, r.frequency) for r in rows][2:] == [
        (date(2027, 7, 1), date(2027, 7, 31), "monthly"), (date(2027, 8, 1), date(2027, 8, 31), "monthly"),
        (date(2027, 9, 1), date(2027, 9, 30), "monthly")]
    with pytest.raises(Exception):
        core_services.change_entity_settings(new, user=t.mp, reason="x", vat_return_period="weekly")


# --- item 7: the alerts job generates the schedules and treats any filed version as filed ---------------------

def test_item7_alerts_for_an_entity_whose_schedule_nobody_generated(t):
    from tax import jobs
    from tax.models import TaxAlert, VatReturnSchedule

    assert not VatReturnSchedule.objects.filter(entity=t.sub).exists()
    jobs.deadline_alerts(today=date(2027, 3, 5))  # SUB is monthly: January was due on 28 February
    assert VatReturnSchedule.objects.filter(entity=t.sub).exists()
    assert TaxAlert.objects.filter(entity=t.sub, kind="vat_overdue", schedule__period_start=Q1).exists()
    assert not TaxAlert.objects.filter(entity=t.sub, schedule__period_start__lt=Q1).exists()  # G4


def test_item7_a_period_with_a_filed_version_and_an_open_amendment_is_not_overdue(t):
    from tax import jobs
    from tax.conftest import journal

    sale(t)
    close_q1(t)
    q1 = filed(t, approved(t))
    entry = journal(t, date(2027, 5, 2), [{"account": "1100", "debit": "20000.00"}, {"account": "2200", "credit": "20000.00"}])
    services.record_correction(q1, user=t.a1, discovered_on=date(2027, 5, 2), description_en="Missed",
                               description_ar="فاتورة لم تُدرج", tax_difference_sar="20000.00", correcting_entries=[entry])
    services.create_amendment(q1, user=t.a1)  # v1 in draft, v0 filed
    kinds = {a.kind for a in jobs.deadline_alerts(today=date(2027, 5, 3))
             if a.entity_id == t.alm.pk and a.schedule_id and a.schedule.period_start == Q1}
    assert "vat_overdue" not in kinds and "vat_payment_overdue" in kinds
    notify = [a for a in jobs.deadline_alerts(today=date(2027, 5, 7)) if a.correction_id]  # 15 days before 22 May
    assert notify[0].incorrect_return_sar == D("5000.00") and "25%" in notify[0].message_en  # C17


# --- item 8: no formulas in exports -------------------------------------------------------------------------------

def test_item8_export_text_is_never_a_formula(t):
    import io

    from openpyxl import load_workbook

    from core.models import Entity

    Entity.objects.filter(pk=t.alm.pk).update(name_en='=HYPERLINK("http://x","y")', name_ar="@SUM(A1)")
    bill(t, [("5300", "1", "10000.00")])  # a claim: negative amounts stay numbers
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    _, data = services.export_return(ret, user=t.v1, fmt="xlsx")
    ws = load_workbook(io.BytesIO(data)).active
    assert (ws.cell(row=1, column=3).data_type, ws.cell(row=1, column=3).value) == ("s", '\'=HYPERLINK("http://x","y")')
    assert ws.cell(row=1, column=4).value == "'@SUM(A1)"
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)
    _, csvdata = services.export_return(ret, user=t.v1, fmt="csv")
    text = csvdata.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and ",=HYPERLINK" not in text and "-1500.00" in text and "'-1500.00" not in text
    assert services._safe_text("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0" and services._safe_text("\tx") == "'\tx"
    assert services._safe_text("-12.50") == "-12.50"


# --- item 9: model guards ------------------------------------------------------------------------------------------

def test_item9_models_refuse_a_non_draft_insert_and_a_filing_without_its_fields(t):
    from tax.models import TaxError

    sale(t)
    draft = services.compute_draft(t.alm, Q1, user=t.a1)
    copy = VatReturn(entity=t.alm, schedule=draft.schedule, period_start=Q1, period_end=draft.period_end,
                     frequency="quarterly", version=1, amends=draft, due_date=draft.due_date, status="filed",
                     created_by=t.a1, created_at=draft.created_at, net_vat_sar=D("-50000"))
    copy._service_write = True
    with pytest.raises(TaxError) as exc:
        copy.save()
    assert exc.value.code == "insert_not_draft"
    close_q1(t)
    from tax.conftest import reviewed

    ret = services.approve_return(reviewed(t, services.submit_for_review(draft, user=t.a1)), user=t.c1)
    ret.status = "filed"
    ret._service_write = True
    with pytest.raises(TaxError) as exc:
        ret.save()
    assert exc.value.code == "filing_incomplete"


def test_item9_a_correction_carried_by_an_approved_return_is_locked(t):
    from tax.conftest import journal, reviewed, prepared
    from tax.models import TaxError, VatCorrection

    old = _filed_q4_2026(t)
    entry = journal(t, date(2027, 2, 2), [{"account": "1100", "debit": "100.00"}, {"account": "2200", "credit": "100.00"}])
    corr = services.record_correction(old, user=t.a1, discovered_on=date(2027, 2, 1), description_en="x",
                                      description_ar="س", tax_difference_sar="100.00", correcting_entries=[entry])
    ret = prepared(t)
    services.acknowledge_exception(ret, user=t.a2, kind="correcting_entry", reference=f"je:{entry.pk}",
                                   decision="confirmed_correction", reason_ar="قيد التصحيح")
    close_q1(t)
    services.approve_return(reviewed(t, ret), user=t.c1)
    corr = VatCorrection.objects.get(pk=corr.pk)
    corr.tax_difference_sar, corr.target_return, corr.status = D("1.00"), None, "open"
    corr._service_write = True
    with pytest.raises(TaxError) as exc:
        corr.save()
    assert exc.value.code == "correction_locked"
    with pytest.raises(TaxError) as exc:
        services.link_correcting_entries(corr, user=t.a1, entries=[])
    assert exc.value.code == "correction_locked"


# --- item 10: authors of adjustments and refund choices are preparers ------------------------------------------

def test_item10_adjustment_author_and_refund_requester_cannot_review_or_approve(t):
    from tax.conftest import pdf

    sale(t)
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    services.add_manual_adjustment(ret, user=t.a2, box_no=1, column="vat", amount_sar="-1000.00",
                                   reason_ar="تخفيض", evidence=pdf())
    ret = services.recompute_draft(ret, user=t.a1)
    ret = services.submit_for_review(ret, user=t.a1)
    with pytest.raises(access.AccessDenied) as exc:
        services.complete_review(ret, user=t.a2)
    assert exc.value.code == "self_approval"
    assert t.a2.pk in services._preparers(VatReturn.objects.get(pk=ret.pk))


def test_item10_refund_requester_counts_as_preparer(t):
    from core.access.testing import member

    bill(t, [("5300", "1", "10000.00")])
    member(t.c1, t.alm, "accountant")
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    services.set_refund_requested(ret, user=t.c1, requested=True)
    ret = services.recompute_draft(ret, user=t.a1)
    assert ret.refund_requested_by == t.c1 and t.c1.pk in services._preparers(ret)
    ret = services.submit_for_review(ret, user=t.a1)
    ret = services.complete_review(ret, user=t.a2)
    close_q1(t)
    with pytest.raises(access.AccessDenied) as exc:
        services.approve_return(ret, user=t.c1)
    assert exc.value.code == "self_approval"


# --- item 11 / C9: vendor credit notes ------------------------------------------------------------------------------

def _vendor_note(t, original, *, vendor_date, accounting_date, price="1000.00", number="VCN-1"):
    from invoicing import services as invoicing

    doc = invoicing.create_purchase_document(t.alm, user=t.ap, vendor=t.vendor, vendor_document_number=number,
                                             vendor_document_date=vendor_date, accounting_date=accounting_date,
                                             doc_type="vendor_credit_note", original_document=original)
    invoicing.add_purchase_line(doc, user=t.ap, account="5300", quantity="1", unit_price=price,
                                description_ar="إشعار دائن")
    doc.refresh_from_db()
    _, net, charged, _, _, _, _ = invoicing._purchase_amounts(doc, list(doc.lines.select_related("vat_rate")))
    invoicing.update_purchase_document(doc, user=t.ap, vendor_total_vat=charged, vendor_total_gross=net + charged)
    invoicing.submit_purchase_document(doc, user=t.ap)
    return invoicing.approve_purchase_document(doc, user=t.c1)


def test_item11_a_reversed_credit_note_goes_to_the_column_of_the_note(t):
    from invoicing import services as invoicing
    from tax.conftest import box

    original = bill(t, [("5300", "1", "10000.00")], on=date(2027, 2, 1))
    close_q1(t)
    filed(t, approved(t), on=date(2027, 4, 20))
    note = _vendor_note(t, original, vendor_date=date(2027, 4, 25), accounting_date=date(2027, 4, 25))
    q2 = services.compute_draft(t.alm, Q2, user=t.a1)
    assert box(q2, 7) == (D("0"), D("-1000.00"), D("-150.00"))  # filed before the note: Adjustment column (R20)
    from invoicing.models import PayableAllocation

    for row in PayableAllocation.objects.filter(vendor_credit_note=note):  # the note was set off against the bill
        invoicing.unallocate(row, user=t.a1)
    invoicing.reverse_document(note, user=t.a1, date=date(2027, 5, 10), reason="Issued in error")
    q2 = services.recompute_draft(q2, user=t.a1)
    assert box(q2, 7) == (D("0"), D("0"), D("0"))  # the reversal undoes the Adjustment column, not the Amount


def test_c9_vendor_credit_note_is_reported_by_the_supplier_note_date(t):
    from tax.conftest import box

    original = bill(t, [("5300", "1", "10000.00")], on=date(2027, 2, 1))
    _vendor_note(t, original, vendor_date=date(2027, 3, 30), accounting_date=date(2027, 4, 2))
    q1 = services.compute_draft(t.alm, Q1, user=t.a1)
    assert box(q1, 7) == (D("9000.00"), D("0"), D("1350.00"))
    tie = services.tie_out(q1, user=t.a1)
    assert tie["difference_input"] == 0 and len(tie["shifted_in"]) == 1
    assert "period_shift" in {e["kind"] for e in services.exceptions(q1, user=t.a1)}
    q2 = services.compute_draft(t.alm, Q2, user=t.a1)
    assert box(q2, 7) == (D("0"), D("0"), D("0"))
    tie = services.tie_out(q2, user=t.a1)
    assert tie["difference_input"] == 0 and len(tie["shifted_out"]) == 1


# --- item 12: rights ----------------------------------------------------------------------------------------------

def test_item12_exceptions_need_ledger_view_like_the_tie_out(t):
    from datetime import timedelta

    from django.utils import timezone

    from core.access.models import Membership, Role, RolePermission
    from core.access.testing import make_user

    role = Role.objects.create(code="tax_reader", name_en="Tax reader", name_ar="قارئ الضرائب", scope="entity",
                               kind="human")
    RolePermission.objects.create(role=role, codename="tax.view")
    reader = make_user("tax_reader")
    Membership.objects.create(user=reader, entity=t.alm, role=role, valid_from=timezone.now() - timedelta(hours=1),
                              granted_by=t.mp, reason="test")
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    assert services.get_return(ret.pk, user=reader) == ret
    for call in (services.exceptions, services.tie_out):
        with pytest.raises(access.AccessDenied):
            call(ret, user=reader)


# --- item 13 ------------------------------------------------------------------------------------------------------

def test_item13_overstatement_time_bar_is_checked_at_approval_and_filing(t):
    from tax.models import TaxError

    sale(t)
    close_q1(t)
    q1 = filed(t, approved(t))
    corr = services.record_correction(q1, user=t.a1, discovered_on=date(2027, 5, 1), description_en="x",
                                      description_ar="س", tax_difference_sar="-100.00")
    q2 = services.compute_draft(t.alm, Q2, user=t.a1)
    corr.refresh_from_db()
    assert corr.target_return_id == q2.pk and services._time_bar(corr) == date(2032, 12, 31)
    services._check_time_bars(q2, date(2032, 12, 31))
    with pytest.raises(TaxError) as exc:
        services._check_time_bars(q2, date(2033, 1, 1))
    assert exc.value.code == "correction_time_barred"


def test_item13_import_cost_account_is_checked_for_sharia_and_restriction(t):
    from ledger.models import Account
    from tax.conftest import pdf
    from tax.models import TaxError

    Account.objects.create(entity=t.alm, code="5950", name_en="Interest cost", name_ar="تكلفة فوائد", type="expense",
                           sharia_prohibited=True)
    for code, expected in (("5950", "debit_account_prohibited"), ("1140", "debit_account_invalid")):
        with pytest.raises(TaxError) as exc:
            services.create_import_vat_record(t.alm, user=t.a1, customs_declaration_no=f"C-{code}",
                                              declaration_date=date(2027, 2, 1), customs_value_sar="100.00",
                                              vat_paid_sar="15.00", attribution="exempt", bank_account="1100",
                                              declaration_file=pdf(), debit_account=code)
        assert exc.value.code == expected


# --- compliance ---------------------------------------------------------------------------------------------------

def _make_legacy(doc):
    """As migration 0026 leaves a document that existed before it: tax point = issue date, basis ``legacy``."""
    from django.db import connection

    from invoicing.models import SalesDocument

    if connection.vendor == "postgresql":
        with connection.cursor() as cursor:
            cursor.execute("SET CONSTRAINTS ALL IMMEDIATE")
            cursor.execute("ALTER TABLE invoicing_salesdocument DISABLE TRIGGER invoicing_salesdocument_lock")
            cursor.execute("UPDATE invoicing_salesdocument SET tax_point_date = issue_date, tax_point_basis = 'legacy' "
                           "WHERE id = %s", [doc.pk])
            cursor.execute("ALTER TABLE invoicing_salesdocument ENABLE TRIGGER invoicing_salesdocument_lock")
            cursor.execute("SET CONSTRAINTS ALL DEFERRED")
    else:
        SalesDocument._base_manager.filter(pk=doc.pk).update(tax_point_date=doc.issue_date, tax_point_basis="legacy")


def test_c10_sales_are_reported_by_tax_point_with_an_issue_date_fallback_for_legacy_rows(t):
    """G6 as built: invoicing stores the tax point, so a new invoice is no fallback; a legacy row (migration 0026)
    is still listed in ``tax_point_fallback`` for the reviewer."""
    from types import SimpleNamespace

    from tax.conftest import prepared

    sale(t)
    ret = prepared(t)
    assert not [e for e in services.exceptions(ret, user=t.a2) if e["kind"] == "tax_point_fallback"]
    services.return_to_draft(ret, user=t.a2, reason_ar="إعادة")
    _make_legacy(sale(t, on=date(2027, 1, 12)))
    ret = services.submit_for_review(services.recompute_draft(ret, user=t.a1), user=t.a1)
    found = [e for e in services.exceptions(ret, user=t.a2) if e["kind"] == "tax_point_fallback"]
    assert len(found) == 1 and found[0]["reference"].startswith("tp:") and len(found[0]["documents"]) == 1
    with pytest.raises(services.TaxError) as exc:
        services.acknowledge_exception(ret, user=t.a2, kind="tax_point_fallback", reference=found[0]["reference"],
                                       decision="not_return_relevant", reason_ar="سبب")
    assert exc.value.code == "decision_invalid"
    invoice = SimpleNamespace(doc_type="invoice", issue_date=date(2027, 4, 2), tax_point_date=date(2027, 3, 28),
                              tax_point_basis="supply")
    assert services._tax_point(invoice) == (date(2027, 3, 28), True)
    legacy = SimpleNamespace(doc_type="invoice", issue_date=date(2027, 4, 2), tax_point_date=date(2027, 4, 2),
                             tax_point_basis="legacy")
    assert services._tax_point(legacy) == (date(2027, 4, 2), False)
    assert services._tax_point(SimpleNamespace(doc_type="invoice", issue_date=date(2027, 4, 2), tax_point_date=None,
                                               tax_point_basis="")) == (date(2027, 4, 2), False)
    note = SimpleNamespace(doc_type="credit_note", issue_date=date(2027, 4, 2), tax_point_date=date(2027, 1, 1),
                           tax_point_basis="issue")
    assert services._tax_point(note) == (date(2027, 4, 2), True)  # notes: their own date (X9)


def test_g6_an_invoice_issued_after_the_quarter_counts_in_the_quarter_of_its_supply(t):
    from tax.conftest import box

    doc = sale(t, on=date(2027, 4, 2), supply_date=date(2027, 3, 28))
    assert (doc.tax_point_date, doc.tax_point_basis) == (date(2027, 3, 28), "supply")
    q1 = services.compute_draft(t.alm, Q1, user=t.a1)
    assert box(q1, 1) == (D("10000.00"), D("0"), D("1500.00"))
    tie = services.tie_out(q1, user=t.a1)
    assert tie["difference_output"] == 0 and tie["shifted_in"] == [doc.journal_entry_id]
    kinds = {e["kind"] for e in services.exceptions(q1, user=t.a1)}
    assert "period_shift" in kinds and "tax_point_fallback" not in kinds
    q2 = services.compute_draft(t.alm, Q2, user=t.a1)
    assert box(q2, 1) == (D("0"), D("0"), D("0"))
    tie = services.tie_out(q2, user=t.a1)
    assert tie["difference_output"] == 0 and tie["shifted_out"] == [doc.journal_entry_id]


def test_g7_box_1_uses_the_sama_tax_point_rate_and_the_tie_out_explains_the_booking_rate(t):
    """invoicing AT-138: SAMA 3.7500 on the supply date, 3.7510 on the issue date 21 days later."""
    from invoicing.conftest import usd_rate
    from tax.conftest import box

    usd_rate(t, "3.7500", on=date(2027, 3, 10))
    usd_rate(t, "3.7510", on=date(2027, 3, 31))
    doc = sale(t, [("rent_commercial", "1", "1000.00")], on=date(2027, 3, 31), currency="USD",
               supply_date=date(2027, 3, 10))
    assert (doc.sar_rate, doc.fx_rate) == (D("3.7500"), D("3.7510"))
    q1 = services.compute_draft(t.alm, Q1, user=t.a1)
    assert box(q1, 1) == (D("3750.00"), D("0"), D("562.50"))  # at the SAMA tax-point rate (C11, G7)
    tie = services.tie_out(q1, user=t.a1)
    assert tie["rate_difference_output"] == D("0.15")  # 2200 booked 562.65
    assert tie["difference_output"] == 0


def test_c11_a_foreign_sale_without_a_sar_rate_is_an_exception_not_the_booking_rate(t):
    from invoicing.conftest import usd_rate
    from invoicing.models import SalesDocument

    usd_rate(t, "3.75")
    doc = sale(t, [("rent_commercial", "1", "1000.00")], currency="USD")
    stored = SalesDocument.objects.select_related("original_document").get(pk=doc.pk)
    stored.sar_rate = None  # in memory only: what the return sees for a document without a stored SAR rate
    comp = services.Computation()
    services._sales_sources(comp, t.alm, stored.journal_entry, stored, 1)
    assert comp.boxes[1] == {"amount": D("0"), "adjustment": D("0"), "vat": D("562.50")}  # no fx_rate amount
    assert [f["kind"] for f in comp.flags] == ["sar_rate_missing"]
    stored.sar_rate = D("3.75")
    comp = services.Computation()
    services._sales_sources(comp, t.alm, stored.journal_entry, stored, 1)
    assert comp.boxes[1]["amount"] == D("3750.00") and comp.flags == []


def _usd_bill(t, vendor_vat_sar=None):
    from invoicing import services as invoicing
    from invoicing.conftest import usd_rate

    usd_rate(t, "3.75", on=date(2027, 1, 15))
    doc = invoicing.create_purchase_document(t.alm, user=t.ap, vendor=t.vendor, vendor_document_number="USD-1",
                                             vendor_document_date=date(2027, 1, 15), accounting_date=date(2027, 1, 15),
                                             currency="USD")
    invoicing.add_purchase_line(doc, user=t.ap, account="5300", quantity="1", unit_price="1000.00",
                                description_ar="بند")
    invoicing.update_purchase_document(doc, user=t.ap, vendor_total_vat="150.00", vendor_total_gross="1150.00",
                                       vendor_vat_sar=vendor_vat_sar)
    invoicing.submit_purchase_document(doc, user=t.ap)
    return invoicing.approve_purchase_document(doc, user=t.c1)


def test_c12_box7_uses_the_suppliers_sar_vat_and_the_rate_difference_stays_in_the_tie_out(t):
    from tax.conftest import box

    _usd_bill(t, vendor_vat_sar="563.00")  # the supplier's SAR VAT; booked at 3.75 = 562.50
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    assert box(ret, 7) == (D("3750.00"), D("0"), D("563.00"))
    tie = services.tie_out(ret, user=t.a1)
    assert (tie["rate_difference_input"], tie["difference_input"]) == (D("-0.50"), 0)


def test_c12_foreign_bill_without_the_sar_vat_is_listed(t):
    _usd_bill(t)
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    assert ret.boxes.get(box_no=7).vat_sar == D("562.50")
    assert "purchase_sar_vat_missing" in {e["kind"] for e in services.exceptions(ret, user=t.a1)}


def test_c13_box8_amount_is_the_import_vat_base(t):
    from tax.conftest import box
    from tax.test_vat_boxes import import_record

    rec = services.create_import_vat_record(
        t.alm, user=t.a1, customs_declaration_no="CUS-B", declaration_date=date(2027, 2, 10),
        customs_value_sar="10000.00", customs_duty_sar="500.00", other_charges_sar="250.00", vat_paid_sar="1612.50",
        attribution="taxable", bank_account="1100", declaration_file=__import__("tax.conftest").conftest.pdf())
    services.submit_tax_document(rec, user=t.a1)
    services.approve_tax_document(rec, user=t.c1)
    assert box(services.compute_draft(t.alm, Q1, user=t.a1), 8) == (D("10750.00"), D("0"), D("1612.50"))
    assert import_record  # the AT-22 helper still builds records without duties


def test_c15_bill_unpaid_12_months_later_is_listed_not_reversed(t):
    bill(t, [("5300", "1", "1000.00")], on=date(2027, 1, 15))
    t.today.set(date(2028, 4, 10))
    services.generate_schedule(t.alm, user=t.a1, until=date(2028, 3, 31))
    ret = services.compute_draft(t.alm, date(2028, 1, 1), user=t.a1)
    found = [e for e in services.exceptions(ret, user=t.a1) if e["kind"] == "unpaid_bill_12_months"]
    assert len(found) == 1 and found[0]["input"] == D("150.00")
    assert ret.boxes.get(box_no=7).vat_sar == 0  # no automatic reversal
    q4 = services.compute_draft(t.alm, date(2027, 10, 1), user=t.a1)
    assert not [e for e in services.exceptions(q4, user=t.a1) if e["kind"] == "unpaid_bill_12_months"]


def test_c16_late_payment_months_are_calendar_months_from_the_due_date():
    due = date(2027, 4, 30)
    assert [services._months_late(due, d) for d in (date(2027, 4, 30), date(2027, 5, 1), date(2027, 5, 31),
                                                    date(2027, 6, 1), date(2027, 6, 30), date(2027, 7, 1))] == \
        [0, 1, 1, 2, 2, 3]
    assert services._months_late(date(2027, 2, 15), date(2027, 3, 15)) == 1
    assert services._months_late(date(2027, 2, 15), date(2027, 3, 16)) == 2
