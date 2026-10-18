"""VAT return lifecycle, entity separation, schedule, alerts and exports (tax-returns.md R1 to R14, R33, R53;
AT-1 to AT-14, AT-30, AT-44)."""
import csv
import io
from datetime import date, timedelta

import pytest
from django.utils import timezone

from core import access
from core.access.models import AccessLog, Membership, Role, RolePermission
from core.access.testing import make_agent, make_user, member
from invoicing import services as invoicing
from invoicing.conftest import make_entity
from tax import jobs, services
from tax.conftest import D, approved, bill, box, close, close_q1, filed, pdf, prepared, reviewed, sale
from tax.models import TaxAlert, TaxError, VatBoxDefinition, VatReturn, VatReturnBox, VatReturnSchedule, VatReturnSourceLine

pytestmark = pytest.mark.django_db
Q1, Q2 = date(2027, 1, 1), date(2027, 4, 1)


# --- entity separation and set-up (R1 to R4) --------------------------------------------------------------

def test_at1_returns_are_scoped_and_other_entity_refused_and_logged(t):
    mine = services.compute_draft(t.alm, Q1, user=t.a1)
    theirs = services.compute_draft(t.sub, Q1, user=t.sub_a)
    assert list(services.vat_returns(t.a1)) == [mine]
    assert services.get_return(mine.pk, user=t.a1) == mine
    with pytest.raises(access.AccessDenied):
        services.get_return(theirs.pk, user=t.a1)
    assert AccessLog.objects.filter(event="permission_denied", actor=t.a1, entity=t.sub, permission="tax.view").exists()
    with pytest.raises(access.AccessDenied):
        services.vat_returns(t.a1, entity=t.sub)
    with pytest.raises(access.AccessDenied):
        services.compute_draft(t.sub, Q1, user=t.a1)


def test_at3_entity_without_vat_number_has_no_vat_return(t):
    with pytest.raises(TaxError) as exc:
        services.compute_draft(t.nov, Q1, user=t.nov_a)
    assert exc.value.code == "not_vat_registered" and "غير مسجل" in str(exc.value)
    assert not VatReturn.objects.filter(entity=t.nov).exists()


def test_r2_vat_returns_need_a_sar_functional_entity(t):
    from core.models import Entity

    Entity.objects.filter(pk=t.sub.pk).update(base_currency="USD")
    with pytest.raises(TaxError) as exc:
        services.compute_draft(t.sub, Q1, user=t.sub_a)
    assert exc.value.code == "vat_requires_sar_functional"


def test_at4_schedule_partial_first_period_frequency_change_and_no_overlap(t):
    new = make_entity("NEW", "full", t.alm, vat="333333333300003", cr="1010000009", vat_effective_date=date(2027, 2, 10))
    acc = make_user("new_acc")
    member(acc, new, "accountant")
    t.today.set(date(2027, 3, 1))
    rows = services.generate_schedule(new, user=acc, until=date(2027, 6, 30))
    assert [(r.period_start, r.period_end, r.due_date) for r in rows] == [
        (date(2027, 2, 10), date(2027, 3, 31), date(2027, 4, 30)), (date(2027, 4, 1), date(2027, 6, 30), date(2027, 7, 31))]
    type(new).objects.filter(pk=new.pk).update(vat_return_period="monthly")  # switched on 2027-05-15 (R4)
    t.today.set(date(2027, 5, 15))
    rows = services.generate_schedule(new, user=acc, until=date(2027, 8, 31))
    assert [(r.period_start, r.period_end, r.frequency) for r in rows][1:] == [
        (date(2027, 4, 1), date(2027, 6, 30), "quarterly"), (date(2027, 7, 1), date(2027, 7, 31), "monthly"),
        (date(2027, 8, 1), date(2027, 8, 31), "monthly")]
    overlap = VatReturnSchedule(entity=new, period_start=date(2027, 7, 15), period_end=date(2027, 9, 30),
                                due_date=date(2027, 10, 31), frequency="quarterly")
    overlap._service_write = True
    with pytest.raises(TaxError) as exc:
        overlap.save()
    assert exc.value.code == "schedule_overlap"


def test_r3_quarters_follow_the_quarter_start_month(t):
    assert services._period_end(date(2027, 2, 1), "quarterly", 2) == date(2027, 4, 30)
    assert services._period_end(date(2027, 12, 1), "quarterly", 1) == date(2027, 12, 31)
    assert services._period_end(date(2027, 12, 1), "quarterly", 3) == date(2028, 2, 29)
    assert services._due_date(date(2027, 12, 31)) == date(2028, 1, 31)


def test_change_tax_settings_needs_the_approver_and_is_logged(t):
    with pytest.raises(access.AccessDenied):
        services.change_tax_settings(t.alm, user=t.a1, vat_correction_threshold_sar="10000.00")
    changed = services.change_tax_settings(t.alm, user=t.c1, alert_days=[10, 3], vat_correction_threshold_sar="10000.00")
    assert (changed.alert_days, changed.vat_correction_threshold_sar) == ([10, 3], D("10000.00"))
    assert AccessLog.objects.filter(event="entity_settings_changed", actor=t.c1, entity=t.alm).exists()
    with pytest.raises(TaxError):
        services.change_tax_settings(t.alm, user=t.c1, fiscal_year_end="06-30")


def test_at5_quarterly_entity_above_40_million_is_warned(t):
    for start, amount in ((date(2026, 4, 1), "13330000.00"), (date(2026, 7, 1), "13330000.00"),
                          (date(2026, 10, 1), "13330000.01")):
        services.import_filed_return(t.alm, user=t.a1, period_start=start, zatca_filing_reference=f"OLD-{start}",
                                     filed_on=start + timedelta(days=100), boxes={1: {"amount": amount}})
    sale(t)
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    assert any("40,000,000" in w and "شهرية" in w for w in services.warnings(ret, user=t.c1))
    assert VatReturn.objects.get(pk=ret.pk).frequency == "quarterly"


def test_at6_unverified_definitions_are_marked_and_approval_allowed(t):
    close_q1(t)
    ret = approved(t)
    assert all(not b.verified for b in ret.boxes.all())
    _, body = services.export_return(ret, user=t.v1, fmt="csv")
    assert "unverified / غير متحقق منه" in body.decode("utf-8-sig")
    definition = VatBoxDefinition.objects.get(box_no=1)
    with pytest.raises(access.AccessDenied):
        services.verify_box_definition(definition, user=t.a1, entity=t.alm)
    assert services.verify_box_definition(definition, user=t.c1, entity=t.alm).verified_at is not None


# --- lifecycle and segregation (R7 to R10, R14) -------------------------------------------------------------

def test_at8_preparer_reviewer_approver_are_three_people(t):
    sale(t)
    close_q1(t)
    ret = prepared(t)
    with pytest.raises(access.AccessDenied) as exc:
        services.complete_review(ret, user=t.a1)
    assert exc.value.code == "self_approval"
    ret = services.complete_review(ret, user=t.c1)
    with pytest.raises(access.AccessDenied) as exc:
        services.approve_return(ret, user=t.c1)
    assert exc.value.code == "self_approval"
    with pytest.raises(access.AccessDenied) as exc:
        services.approve_return(ret, user=t.a2)
    assert exc.value.code == "no_membership"
    services.return_to_draft(ret, user=t.c1, reason_ar="مراجعة من شخص آخر", reason_en="Second reviewer")
    ret = services.approve_return(reviewed(t, prepared(t)), user=t.c1)
    assert (ret.status, ret.submitted_by, ret.reviewed_by, ret.approved_by) == ("approved", t.a1, t.a2, t.c1)


def test_r8_preparer_holding_approver_rights_cannot_approve(t):
    member(t.a1, t.alm, "approver")
    close_q1(t)
    ret = reviewed(t, prepared(t))
    with pytest.raises(access.AccessDenied) as exc:
        services.approve_return(ret, user=t.a1)
    assert exc.value.code == "self_approval"


def test_at9_open_periods_block_approval(t):
    close(t, "2027-01", "2027-02")
    ret = reviewed(t, prepared(t))
    with pytest.raises(TaxError) as exc:
        services.approve_return(ret, user=t.c1)
    assert exc.value.code == "periods_open" and "2027-03" in str(exc.value)


def test_at10_changed_sources_send_the_return_back_to_draft(t):
    doc = bill(t, [("5300", "1", "1000.00")], on=date(2027, 3, 10))
    ret = reviewed(t, prepared(t))
    invoicing.reverse_document(doc, user=t.a1, date=date(2027, 3, 20), reason="wrong vendor")
    close_q1(t)
    with pytest.raises(TaxError) as exc:
        services.approve_return(ret, user=t.c1)
    assert exc.value.code == "sources_changed"
    stored = VatReturn.objects.get(pk=ret.pk)
    assert stored.status == "draft" and stored.returned_reason_ar and stored.reviewed_by is None


def test_r10_review_also_refuses_changed_sources(t):
    ret = prepared(t)
    sale(t)
    with pytest.raises(TaxError) as exc:
        services.complete_review(ret, user=t.a2)
    assert exc.value.code == "sources_changed"


def test_at7_approved_return_is_locked(t):
    sale(t)
    close_q1(t)
    ret = approved(t)
    assert len(ret.source_hash) == 64
    with pytest.raises(TaxError) as exc:
        services.compute_draft(t.alm, Q1, user=t.a1)
    assert exc.value.code == "return_locked"
    for call in (lambda: services.recompute_draft(ret, user=t.a1),
                 lambda: services.submit_for_review(ret, user=t.a1),
                 lambda: services.set_refund_requested(ret, user=t.a1, requested=True)):
        with pytest.raises(TaxError):
            call()
    b = ret.boxes.get(box_no=1)
    b.vat_sar = D("1.00")
    b._service_write = True
    with pytest.raises(TaxError) as exc:
        b.save()
    assert exc.value.code == "return_locked"
    stored = VatReturn.objects.get(pk=ret.pk)
    stored.net_vat_sar = D("0")
    stored._service_write = True
    with pytest.raises(TaxError):
        stored.save()
    with pytest.raises(TaxError):
        VatReturn.objects.filter(pk=ret.pk).update(net_vat_sar=0)
    with pytest.raises(TaxError):
        VatReturn.objects.get(pk=ret.pk).delete()
    line = VatReturnSourceLine.objects.filter(vat_return=ret).first()
    line._service_write = True
    with pytest.raises(TaxError):
        line.delete()
    assert box(VatReturn.objects.get(pk=ret.pk), 1) == (D("10000.00"), D("0"), D("1500.00"))


def test_models_refuse_writes_outside_the_services(t):
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    with pytest.raises(TaxError) as exc:
        VatReturnBox.objects.filter(vat_return=ret).update(vat_sar=1)
    assert exc.value.code == "service_only"
    stored = VatReturn.objects.get(pk=ret.pk)
    stored.status = "approved"
    with pytest.raises(TaxError):
        stored.save()


def test_at11_drill_down_sums_to_the_box_and_follows_rights(t):
    first, second = sale(t), sale(t, [("rent_commercial", "1", "2000.00")], on=date(2027, 2, 10))
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    rows = services.drill_down(ret, user=t.v1, box_no=1, column="vat")
    assert sum(r["amount_sar"] for r in rows) == box(ret, 1)[2] == D("1800.00")
    assert {r["journal_entry"] for r in rows} == {first.journal_entry_id, second.journal_entry_id}
    assert {r["source_id"] for r in rows} == {first.pk, second.pk}
    with pytest.raises(access.AccessDenied):
        services.drill_down(ret, user=t.sub_a, box_no=1)


def test_at14_agent_drafts_only(t):
    role = Role.objects.create(code="agent_tax_drafter", name_en="Agent: VAT drafter", name_ar="وكيل إعداد الإقرار",
                               scope="entity", kind="agent")
    for codename in ("tax.view", "tax.prepare_return"):
        RolePermission.objects.create(role=role, codename=codename)
    agent = make_agent("tax_drafter")
    Membership.objects.create(user=agent, entity=t.alm, role=role, valid_from=timezone.now() - timedelta(hours=1),
                              granted_by=t.mp, reason="pilot")
    sale(t)
    ret = services.compute_draft(t.alm, Q1, user=agent)
    assert (ret.source, ret.created_by, box(ret, 1)[2]) == ("agent", agent, D("1500.00"))
    assert AccessLog.objects.filter(actor=agent, details__action="compute_draft").exists()
    with pytest.raises(access.AccessDenied) as exc:
        services.submit_for_review(ret, user=agent)
    assert exc.value.code == "agent_forbidden"
    with pytest.raises(access.AccessDenied):
        services.add_manual_adjustment(ret, user=agent, box_no=1, column="vat", amount_sar="1.00", reason_ar="س",
                                       evidence=pdf())
    ret = services.submit_for_review(ret, user=t.a1)  # a person takes it over
    assert ret.submitted_by == t.a1


# --- posting T1 (R33, AT-30) --------------------------------------------------------------------------------

def lines_of(entry):
    return sorted((l.account.code, l.debit, l.credit) for l in entry.lines.select_related("account"))


def test_at30_offset_entry_at_approval(t):
    sale(t)
    bill(t, [("5300", "1", "4000.00")])
    close_q1(t)
    ret = approved(t)
    entry = ret.settlement_entry
    assert lines_of(entry) == [("1310", 0, D("600.00")), ("2200", D("600.00"), 0)]
    assert (entry.date, entry.posted_by, entry.document_type, entry.document_id) == (
        date(2027, 4, 10), t.c1, "tax.vatreturn", ret.pk)
    assert box(ret, 16)[2] == D("900.00")


def test_at30_claim_offsets_output_only_and_zero_posts_nothing(t):
    sale(t, [("rent_commercial", "1", "1000.00")])
    bill(t, [("5300", "1", "10000.00")])
    close_q1(t)
    ret = approved(t)
    assert lines_of(ret.settlement_entry) == [("1310", 0, D("150.00")), ("2200", D("150.00"), 0)]
    assert box(ret, 16)[2] == D("-1350.00")
    close(t, "2027-04", "2027-05", "2027-06")
    t.today.set(date(2027, 7, 10))
    filed(t, ret)
    q2 = approved(t, Q2)
    assert q2.settlement_entry is None and q2.settlement_offset_sar == 0 and box(q2, 15)[2] == D("1350.00")
    assert box(q2, 16)[2] == D("-1350.00")


def test_r33_t1_posts_only_through_the_service(t):
    from ledger import documents

    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    with pytest.raises(TaxError) as exc:
        documents.get("tax.vatreturn").build_lines(ret)
    assert exc.value.code == "post_through_services"


def test_earlier_open_return_blocks_approval(t):
    services.compute_draft(t.alm, Q1, user=t.a1)
    close_q1(t)
    close(t, "2027-04", "2027-05", "2027-06")
    t.today.set(date(2027, 7, 10))
    q2 = reviewed(t, prepared(t, Q2))
    with pytest.raises(TaxError) as exc:
        services.approve_return(q2, user=t.c1)
    assert exc.value.code == "earlier_return_open"


# --- filing and imports (R12) -------------------------------------------------------------------------------

def test_filing_needs_reference_date_and_acknowledgement(t):
    close_q1(t)
    ret = approved(t)
    with pytest.raises(TaxError) as exc:
        services.record_filing(ret, user=t.a1, zatca_filing_reference="", filed_on=date(2027, 4, 20),
                               acknowledgement=pdf())
    assert exc.value.code == "filing_incomplete"
    with pytest.raises(access.AccessDenied):
        services.record_filing(ret, user=t.v1, zatca_filing_reference="Z-1", filed_on=date(2027, 4, 20),
                               acknowledgement=pdf())
    done = filed(t, ret)
    assert (done.status, done.zatca_filing_reference, done.filed_by) == ("settled", "ZATCA-TEST-1", t.a1)  # nothing due


def test_at12_imported_legacy_return_and_box15(t):
    old = services.import_filed_return(t.alm, user=t.a1, period_start=date(2026, 10, 1),
                                       zatca_filing_reference="FALCON-2026Q4", filed_on=date(2027, 1, 25),
                                       boxes={1: {"amount": "20000.00", "vat": "3000.00"},
                                              7: {"amount": "46666.67", "vat": "7000.00"},
                                              13: {"vat": "-4000.00"}, 16: {"vat": "-4000.00"}})
    assert (old.status, old.source, old.net_vat_sar, old.source_lines.count()) == ("filed", "import", D("-4000.00"), 0)
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    assert box(ret, 15)[2] == D("4000.00") and box(ret, 16)[2] == D("-4000.00")
    with pytest.raises(TaxError) as exc:
        services.import_filed_return(t.alm, user=t.a1, period_start=date(2026, 7, 1), zatca_filing_reference="X",
                                     filed_on=date(2026, 10, 20), boxes={16: {"vat": "5.00"}})
    assert exc.value.code == "import_boxes_inconsistent"


def test_r25_refund_requested_claim_is_not_carried(t):
    services.import_filed_return(t.alm, user=t.a1, period_start=date(2026, 10, 1), zatca_filing_reference="F",
                                 filed_on=date(2027, 1, 25), refund_requested=True,
                                 boxes={13: {"vat": "-4000.00"}, 16: {"vat": "-4000.00"}})
    assert box(services.compute_draft(t.alm, Q1, user=t.a1), 15)[2] == D("0")


# --- alerts (R13, AT-13) --------------------------------------------------------------------------------------

def test_at13_deadline_alerts(t):
    sale(t)
    services.compute_draft(t.alm, Q1, user=t.a1)
    for day in (date(2027, 4, 14), date(2027, 4, 15), date(2027, 4, 25), date(2027, 4, 29), date(2027, 5, 1)):
        jobs.deadline_alerts(today=day)
    jobs.deadline_alerts(today=date(2027, 5, 1))  # idempotent
    rows = TaxAlert.objects.filter(entity=t.alm, schedule__period_start=Q1).order_by("alert_date")
    assert [(a.alert_date, a.kind) for a in rows] == [
        (date(2027, 4, 15), "vat_due_soon"), (date(2027, 4, 25), "vat_due_soon"), (date(2027, 4, 29), "vat_due_soon"),
        (date(2027, 5, 1), "vat_overdue")]
    overdue = rows.last()
    assert overdue.exposure_sar == D("75.00") and "غير متحقق" in overdue.message_ar  # 1,500.00 x 5% x 1 month
    assert len(services.alerts(t.alm, user=t.a1)) == len(services.alerts(t.alm, user=t.c1)) > 4  # earlier periods too
    # periods since VAT registration (2024-01-01) that have no return yet are overdue as well, until imported (R12)
    assert TaxAlert.objects.filter(entity=t.alm, kind="vat_overdue", schedule__period_start=date(2026, 10, 1)).exists()
    with pytest.raises(access.AccessDenied):
        services.alerts(t.alm, user=t.v1)
    close_q1(t)
    ret = services.approve_return(reviewed(t, prepared(t)), user=t.c1)
    filed(t, ret, on=date(2027, 5, 2))
    payment = services.create_vat_payment(ret, user=t.a1, amount_sar="1500.00", payment_date=date(2027, 5, 3),
                                          bank_account="1100")
    services.submit_tax_document(payment, user=t.a1)
    services.approve_tax_document(payment, user=t.c1)
    assert [a for a in jobs.deadline_alerts(today=date(2027, 5, 5)) if a.schedule.period_start == Q1] == []


def test_r13_payment_overdue_alert_after_filing(t):
    sale(t)
    close_q1(t)
    filed(t, approved(t))
    alerts = [a for a in jobs.deadline_alerts(today=date(2027, 5, 2)) if a.schedule.period_start == Q1]
    assert [a.kind for a in alerts] == ["vat_payment_overdue"] and alerts[0].exposure_sar == D("75.00")


# --- exports (R53, AT-44) -------------------------------------------------------------------------------------

def test_at44_export_csv_and_xlsx_in_both_languages(t):
    from openpyxl import load_workbook

    sale(t)
    ret = services.compute_draft(t.alm, Q1, user=t.a1)
    name, body = services.export_return(ret, user=t.v1, fmt="csv", language="en")
    rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))
    box1 = next(r for r in rows if r and r[0] == "1")
    assert box1[1:5] == ["Standard rated sales", "10000.00", "0.00", "1500.00"]
    assert box1[6] == "المبيعات الخاضعة للنسبة الأساسية" and name.endswith("-en.csv")
    assert any("ZATCA portal" in r[0] and "بوابة الهيئة" in r[0] for r in rows if r)
    name, body = services.export_return(ret, user=t.v1, fmt="xlsx", language="ar")
    ws = load_workbook(io.BytesIO(body)).active
    assert ws.sheet_view.rightToLeft
    values = [[c.value for c in row] for row in ws.iter_rows()]
    header = next(r for r in values if r[0] == "الخانة")
    assert header[1] == "التسمية"
    box1 = next(r for r in values if r[0] == 1)
    assert box1[1] == "المبيعات الخاضعة للنسبة الأساسية" and box1[6] == "Standard rated sales"
    assert D(str(box1[4])) == D("1500.00")
    with pytest.raises(TaxError):
        services.export_return(ret, user=t.v1, fmt="pdf")
    with pytest.raises(access.AccessDenied):
        services.export_return(ret, user=t.sub_a)


def test_at44_every_label_is_bilingual():
    import re

    from tax import reference

    arabic = re.compile("[؀-ۿ]")
    for n, (en, ar) in reference.BOX_LABELS.items():
        assert en and arabic.search(ar) and not arabic.search(en), n
    assert len(reference.BOXES) == 16 and all("[verify]" in row[7] or "confirmed" in row[7]
                                             for row in reference.BOXES.values())


def test_admin_is_read_only_and_entity_scoped(client, t):
    mine = services.compute_draft(t.alm, Q1, user=t.a1)
    theirs = services.compute_draft(t.sub, Q1, user=t.sub_a)
    t.a1.is_staff = True
    t.a1.save()
    client.force_login(t.a1)
    resp = client.get("/admin/tax/vatreturn/")
    assert resp.status_code == 200
    body = resp.content.decode()
    assert f"/admin/tax/vatreturn/{mine.pk}/" in body and f"/admin/tax/vatreturn/{theirs.pk}/" not in body
    assert client.get(f"/admin/tax/vatreturn/{theirs.pk}/change/").status_code in (302, 403, 404)
    assert client.get("/admin/tax/vatreturn/add/").status_code == 403
