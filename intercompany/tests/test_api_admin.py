"""REST API /api/v1/intercompany/, admin scoping and the seeded permissions/approval policies (§7)."""
import datetime
import io
import uuid
from decimal import Decimal

import pytest
from django.contrib.admin.sites import site
from django.test import RequestFactory

from access.catalogue import ROLES, SERVICE_FORBIDDEN
from access.models import Role
from approvals.models import ApprovableType, ApprovalPolicy, ApprovalRequest
from intercompany import services
from intercompany.models import IntercompanyBill, IntercompanyInvoice, ServiceAgreement, ServiceType, VatRate

from .conftest import approve

pytestmark = pytest.mark.django_db
D = Decimal
BASE = "/api/v1/intercompany"

AGREEMENT = {
    "entity": "ALM", "counterparty": "LFN", "service_type": "FIN", "description_en": "Finance services",
    "description_ar": "خدمات مالية", "pricing_method": "fixed_fee", "fixed_fee": "10000.00", "vat_treatment": "standard",
    "billing_frequency": "monthly", "payment_terms_days": 30, "start_date": "2026-01-01",
}


def approve_via_api(client, doc_type, obj_id):
    r = ApprovalRequest.objects.get(document_type__code=doc_type, object_id=str(obj_id), status="pending")
    return client.post(f"/api/v1/approvals/requests/{r.pk}/approve", {"document_version": r.document_version},
                       format="json", HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()))


# ---------------------------------------------------------------- end to end through the API


def test_api_full_cycle_agreement_run_approve_issue_bill(api, accountant, mp, lfn_acc, alm, lfn):
    acc, boss, rec = api(accountant), api(mp), api(lfn_acc)
    r = acc.post(f"{BASE}/agreements", AGREEMENT, format="json")
    assert r.status_code == 201, r.json()
    body = r.json()
    assert body["number"] == "ALM-AGR-2026-0001" and body["fixed_fee"] == "10000.00"
    assert body["description"] == {"en": "Finance services", "ar": "خدمات مالية"} and body["source"] == "human"
    agreement_id = body["id"]
    assert acc.post(f"{BASE}/agreements/{agreement_id}/submit").status_code == 201
    assert approve_via_api(boss, "intercompany.agreement", agreement_id).status_code == 200
    r = acc.post(f"{BASE}/billing-runs", {"entity": "ALM", "period": "2026-11"}, format="json")
    assert r.status_code == 201 and r.json()["invoices_created"] == 1
    inv = acc.get(f"{BASE}/invoices?entity=ALM").json()[0]
    assert (inv["net_total"], inv["vat_total"], inv["gross_total"]) == ("10000.00", "1500.00", "11500.00")
    assert inv["status"] == "draft" and inv["lines"][0]["vat_category"] == "S"
    r = acc.post(f"{BASE}/invoices/{inv['id']}/issue")
    assert r.status_code == 422 and r.json()["code"] == "not_approved" and r.json()["message_ar"]
    assert acc.post(f"{BASE}/invoices/{inv['id']}/submit").status_code == 201
    r = approve_via_api(boss, "intercompany.invoice", inv["id"])
    assert r.status_code == 200 and r.json()["status"] == "approved"
    inv = acc.get(f"{BASE}/invoices/{inv['id']}").json()
    assert inv["status"] == "issued" and inv["number"] == "ALM-ICI-2026-000001" and inv["journal_entry"]
    assert inv["zatca"]["ready"] is True and inv["zatca"]["status"] == "not_submitted"
    assert inv["mirror_bill"] == {"exists": True}  # R19: the accountant has no rights in LFN
    bills = rec.get(f"{BASE}/bills").json()
    assert len(bills) == 1 and bills[0]["gross_total"] == "11500.00" and bills[0]["entity"]["code"] == "LFN"
    bill_id, line_id = bills[0]["id"], bills[0]["lines"][0]["id"]
    r = rec.patch(f"{BASE}/bills/{bill_id}", {"lines": [{"id": line_id, "net_amount": "1"}]}, format="json")
    assert r.status_code == 422
    r = rec.patch(f"{BASE}/bills/{bill_id}", {"lines": [{"id": line_id, "expense_account": "5900"}]}, format="json")
    assert r.status_code == 200 and r.json()["lines"][0]["expense_account"] == "5900"
    assert rec.post(f"{BASE}/bills/{bill_id}/submit").status_code == 201
    assert approve_via_api(boss, "intercompany.bill", bill_id).status_code == 200
    assert rec.get(f"{BASE}/bills/{bill_id}").json()["status"] == "posted"
    rep = boss.get(f"{BASE}/reconciliation?as_of=2026-11-30").json()
    row = rep["rows"][0]
    assert row["issuer_balance"] == row["receiver_balance"] == "11500.00" and row["difference"] == "0.00"
    assert rep["headings"]["issuer"] == {"en": "Issuer", "ar": "المُصدر"}


def test_api_scoping_404_and_403(api, issued_invoice, accountant, lfn_acc, make_user, alm, lfn):
    inv = issued_invoice()
    rec, acc = api(lfn_acc), api(accountant)
    assert rec.get(f"{BASE}/invoices").json() == []
    assert rec.get(f"{BASE}/invoices/{inv.pk}").status_code == 404
    assert rec.post(f"{BASE}/invoices/{inv.pk}/credit-notes", {"reason_en": "x", "reason_ar": "س"},
                    format="json").status_code == 404
    assert acc.get(f"{BASE}/bills").json() == []
    assert acc.get(f"{BASE}/bills/{inv.mirror_bill.pk}").status_code == 404
    analyst = api(make_user("ana", (alm, "business_analyst")))
    assert analyst.get(f"{BASE}/agreements").json() == []
    assert analyst.post(f"{BASE}/agreements", AGREEMENT, format="json").status_code == 403
    assert analyst.get(f"{BASE}/reconciliation?as_of=2026-11-30").status_code == 403
    r = acc.post(f"{BASE}/agreements", {**AGREEMENT, "entity": "LFN"}, format="json")
    assert r.status_code == 404  # no access to LFN at all
    from rest_framework.test import APIClient

    assert APIClient().get(f"{BASE}/invoices").status_code in (401, 403)  # signed in with 2FA only (D14)


def test_api_credit_note_and_cancel(api, issued_invoice, accountant, make_agreement, alm, nov):
    inv = issued_invoice()
    acc = api(accountant)
    r = acc.post(f"{BASE}/invoices/{inv.pk}/credit-notes",
                 {"reason_en": "Partial", "reason_ar": "جزئي", "lines": [{"line": inv.lines.get().pk, "net_amount": "5000"}]},
                 format="json")
    assert r.status_code == 201 and r.json()["gross_total"] == "5750.00" and r.json()["kind"] == "credit_note"
    assert r.json()["credit_of"] == inv.number
    r = acc.post(f"{BASE}/invoices/{inv.pk}/credit-notes", {"reason_en": "More", "reason_ar": "المزيد",
                 "lines": [{"line": inv.lines.get().pk, "net_amount": "5217.39"}]}, format="json")
    assert r.status_code == 422 and r.json()["code"] in ("credit_exceeds_line", "credit_exceeds_balance")
    a2 = make_agreement()
    services.run_billing(accountant, alm, nov)
    draft = a2.invoices.get()
    r = acc.patch(f"{BASE}/invoices/{draft.pk}", {"net_total": "1"}, format="json")
    assert r.status_code == 422
    r = acc.patch(f"{BASE}/invoices/{draft.pk}", {"lines": [{"id": draft.lines.get().pk, "description_en": "Nov fee"}]},
                  format="json")
    assert r.status_code == 200 and r.json()["lines"][0]["description"]["en"] == "Nov fee"
    assert acc.post(f"{BASE}/invoices/{draft.pk}/cancel", {"reason": "dup"}, format="json").json()["status"] == "cancelled_draft"


def test_api_agent_token_drafts_only(api, make_agreement, make_user, alm, nov):
    a = make_agreement()
    bot = make_user("ic-bot", (alm, "agent_drafter"), user_type="service_agent", agent_name="ic-bot")
    c = api(bot)
    r = c.post(f"{BASE}/invoices", {"agreement": a.pk, "period": "2026-11"}, format="json")
    assert r.status_code == 201 and r.json()["source"] == "agent" and r.json()["source_agent"] == "ic-bot"
    assert c.post(f"{BASE}/billing-runs", {"entity": "ALM", "period": "2026-11"}, format="json").status_code == 403
    assert c.post(f"{BASE}/invoices/{r.json()['id']}/issue").status_code == 403
    assert c.get(f"{BASE}/invoices").json()[0]["id"] == r.json()["id"]  # own drafts only


def test_api_timesheets_and_pools(api, make_agreement, accountant, mp, alm):
    a = make_agreement("timesheet")
    acc, boss = api(accountant), api(mp)
    r = acc.post(f"{BASE}/timesheets", {"agreement": a.pk, "date": "2026-11-02", "hours": "3.50",
                                        "description_en": "Close", "description_ar": "إقفال"}, format="json")
    assert r.status_code == 201 and r.json()["hours"] == "3.50"
    tid = r.json()["id"]
    assert acc.patch(f"{BASE}/timesheets/{tid}", {"hours": "4"}, format="json").json()["hours"] == "4.00"
    assert acc.post(f"{BASE}/timesheets/{tid}/submit").json()["status"] == "submitted"
    assert acc.post(f"{BASE}/timesheets/{tid}/approve").status_code == 403  # own hours
    assert boss.post(f"{BASE}/timesheets/{tid}/approve").json()["status"] == "approved"
    r = acc.post(f"{BASE}/cost-pools", {"entity": "ALM", "code": "BAD", "name_en": "Bad", "name_ar": "سيئ",
                                        "accounts": ["5610"]}, format="json")
    assert r.status_code == 422 and r.json()["code"] == "sharia_prohibited"
    r = acc.post(f"{BASE}/cost-pools", {"entity": "ALM", "code": "FIN", "name_en": "Finance", "name_ar": "المالية",
                                        "accounts": ["5100"]}, format="json")
    assert r.status_code == 201 and r.json()["accounts"] == ["5100"]
    r = acc.post(f"{BASE}/cost-pools/{r.json()['id']}/compute", {"period": "2026-11"}, format="json")
    assert r.json() == {"pool": "FIN", "period": "2026-11", "amount": "0.00", "locked": False}
    assert len(acc.get(f"{BASE}/service-types").json()) == ServiceType.objects.count() == 11


def test_api_reconciliation_xlsx_bilingual(api, issued_invoice, accountant):
    from openpyxl import load_workbook

    issued_invoice()
    for lang, first in (("en", "Issuer"), ("ar", "المُصدر")):
        r = api(accountant).get(f"{BASE}/reconciliation?as_of=2026-11-30&export=xlsx&lang={lang}")
        assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"]
        ws = load_workbook(io.BytesIO(r.content)).active
        assert ws.cell(1, 1).value == first and ws.cell(2, 1).value == "ALM" and ws.cell(2, 4).value == 11500
    assert api(accountant).get(f"{BASE}/reconciliation").status_code == 422


# ---------------------------------------------------------------- admin


def _req(user):
    r = RequestFactory().get("/admin/")
    r.user = user
    r.META["REMOTE_ADDR"] = "127.0.0.1"
    return r


def test_admin_scoped_and_read_only(issued_invoice, accountant, lfn_acc, make_user, alm):
    inv = issued_invoice()
    inv_admin, bill_admin = site._registry[IntercompanyInvoice], site._registry[IntercompanyBill]
    assert list(inv_admin.get_queryset(_req(accountant))) == [inv]
    assert list(inv_admin.get_queryset(_req(lfn_acc))) == []
    assert list(bill_admin.get_queryset(_req(accountant))) == []
    assert list(bill_admin.get_queryset(_req(lfn_acc))) == [inv.mirror_bill]
    for admin in (inv_admin, bill_admin, site._registry[ServiceAgreement]):
        assert not admin.has_add_permission(_req(accountant))
        assert not admin.has_change_permission(_req(accountant), None)
        assert not admin.has_delete_permission(_req(accountant), None)
    analyst = make_user("ana", (alm, "business_analyst"))
    assert not inv_admin.has_module_permission(_req(analyst))
    st_admin = site._registry[ServiceType]
    assert not st_admin.has_change_permission(_req(accountant))
    mp_user = make_user("mp2", (alm, "managing_partner"))
    assert st_admin.has_change_permission(_req(mp_user))


def test_admin_changelist_pages_render(api, issued_invoice, accountant):
    issued_invoice()
    accountant.is_staff = True
    accountant.save()
    c = api(accountant)
    for model in ("intercompanyinvoice", "serviceagreement", "billingrun", "timesheetentry", "costpool"):
        assert c.get(f"/admin/intercompany/{model}/").status_code == 200, model
    inv = IntercompanyInvoice.objects.get()
    assert c.get(f"/admin/intercompany/intercompanyinvoice/{inv.pk}/change/").status_code == 200


# ---------------------------------------------------------------- seeded permissions and policies (§7)


def test_role_seeding_matches_spec_permissions():
    ic = lambda role: {p for p in ROLES[role][2] if p.startswith("intercompany.")}  # noqa: E731
    assert "intercompany.billing.run" in ic("accountant") and "intercompany.invoice.issue" in ic("accountant")
    assert {"intercompany.agreement.create", "intercompany.timesheet.create", "intercompany.invoice.create"} == ic("pa_drafter")
    assert ic("pa_drafter") == ic("agent_drafter")
    assert ic("auditor") == {"intercompany.agreement.view", "intercompany.invoice.view", "intercompany.bill.view",
                             "intercompany.report.view"}
    assert ic("business_analyst") == set()
    assert "intercompany.servicetype.manage" in ic("managing_partner")
    assert "intercompany.invoice.issue" not in ic("cfo")
    for p in ("intercompany.billing.run", "intercompany.invoice.issue", "intercompany.timesheet.approve",
              "intercompany.bill.edit"):
        assert p in SERVICE_FORBIDDEN
    db = set(Role.objects.get(code="accountant").permissions.filter(code__startswith="intercompany.")
             .values_list("code", flat=True))
    assert db == ic("accountant")  # the data migration seeded the database the same way


def test_seeded_approval_types_route_to_managing_partner_and_vat_rate():
    for code in ("intercompany.agreement", "intercompany.invoice", "intercompany.bill"):
        t = ApprovableType.objects.get(code=code)
        p = ApprovalPolicy.objects.get(document_type=t, entity__isnull=True, currency="SAR")
        assert [s.approver_role.code for s in p.steps.all()] == ["managing_partner"]
    assert VatRate.objects.get(country="SA").rate == D("15.00")
    assert ServiceType.objects.get(code="RENT_VEH").income_account_code == "4110"
    assert ServiceType.objects.get(code="SECOND").expense_account_code == "5100"
